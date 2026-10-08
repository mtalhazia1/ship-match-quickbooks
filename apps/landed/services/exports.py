"""CSV and Excel files of landed cost: one shipment, or the per-product report.

CSV rows go through apps.core.csvsafe so a product name like "=HYPERLINK(...)" stays text in Excel. In
the .xlsx files every text cell is stored as text (never as a formula) for the same reason, and only
the amount columns are stored as numbers (a SKU like 00123 keeps its zeros).
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal

from apps.core import csvsafe

from ..models import Category

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONEY, UNIT, QTY, PCT = "#,##0.00", "#,##0.0000", "#,##0.####", "0.0"


@dataclass
class Sheet:
    name: str
    head: list[str]
    rows: list[list]
    formats: dict = field(default_factory=dict)   # column index -> number format
    title: str = ""


def _n(value, places: int = 2) -> str:
    if value is None or value == "":
        return ""
    return f"{Decimal(value):.{places}f}"


def _qty(value) -> str:
    if value is None:
        return ""
    text = f"{Decimal(value):.4f}"
    return text.rstrip("0").rstrip(".") if "." in text else text


def shipment_sheet(lc) -> Sheet:
    cats = [c for c, _ in Category.choices]
    cur = lc.currency
    head = ["SKU", "Description", "HS code", "Supplier", "Quantity", "Weight (kg)", "Volume (cbm)",
            f"Goods value ({cur})", *[f"{Category(c).label} ({cur})" for c in cats], f"Charges ({cur})",
            f"Landed cost ({cur})", f"Landed cost per unit ({cur})", "Uplift (%)"]
    rows = [[p.sku, p.description, p.hs_code, p.vendor_name, _qty(p.quantity), _qty(p.weight_kg),
             _qty(p.volume_cbm), _n(p.value_home), *[_n(p.part(c)) for c in cats], _n(p.charges), _n(p.landed),
             _n(p.per_unit, 4), _n(p.uplift, 1)] for p in lc.products]
    if lc.products and lc.complete:
        rows.append(["", "Total", "", "", "", "", "", _n(lc.goods_total), *[_n(lc.category_total(c)) for c in cats],
                     _n(lc.charges_total), _n(lc.landed_total), "", _n(lc.uplift, 1)])
    money_cols = range(7, 7 + len(cats) + 3)
    formats = {4: QTY, 5: QTY, 6: QTY, **{i: MONEY for i in money_cols}, len(head) - 2: UNIT, len(head) - 1: PCT}
    return Sheet("Landed cost", head, rows, formats, title=f"Landed cost of {lc.shipment.reference} in {cur}")


def charges_sheet(lc) -> Sheet:
    head = ["Source", "Charge", "Type", "Amount", "Currency", f"Amount ({lc.currency})", "Spread by"]
    rows = [[r.charge.source, r.charge.description, Category(r.charge.category).label, _n(r.charge.amount),
             r.charge.currency, _n(r.amount_home), r.used or r.basis] for r in lc.charges]
    return Sheet("Charges", head, rows, {3: MONEY, 5: MONEY})


def shipment_csv(lc) -> bytes:
    s = shipment_sheet(lc)
    return _csv([s.head, *s.rows])


def shipment_xlsx(lc) -> bytes:
    notes = Sheet("Notes", ["Note"], [[n.text] for n in lc.notes] or [["No notes."]])
    return _xlsx([shipment_sheet(lc), charges_sheet(lc), notes])


def report_sheets(report) -> tuple[Sheet, Sheet]:
    cur = report.currency
    summary = Sheet("By product", ["SKU", "Description", "Supplier", "Shipments", "Quantity",
                                   f"Average landed cost per unit ({cur})", f"Last landed cost per unit ({cur})",
                                   "Last shipment date", "Change from previous (%)", f"Landed cost ({cur})",
                                   "Average uplift (%)"],
                    [[r.sku, r.description, r.vendor_name, len(r.points), _qty(r.quantity), _n(r.average, 4),
                      _n(r.last.per_unit, 4) if r.last else "", r.last.as_of.isoformat() if r.last else "",
                      _n(r.change, 1), _n(r.landed_total), _n(r.uplift, 1)] for r in report.rows],
                    {3: "0", 4: QTY, 5: UNIT, 6: UNIT, 8: PCT, 9: MONEY, 10: PCT},
                    title=f"Landed cost by product, {report.start:%d %b %Y} to {report.end:%d %b %Y}, in {cur}")
    detail = Sheet("By shipment", ["SKU", "Description", "Supplier", "Date", "Shipment", "Status", "Quantity",
                                   f"Goods value ({cur})", f"Charges ({cur})", f"Landed cost ({cur})",
                                   f"Per unit ({cur})"],
                   [[r.sku, r.description, r.vendor_name, pt.as_of.isoformat(), pt.reference,
                     "In review" if pt.in_review else "Approved", _qty(pt.quantity), _n(pt.goods), _n(pt.charges),
                     _n(pt.landed), _n(pt.per_unit, 4)] for r in report.rows for pt in r.points],
                   {6: QTY, 7: MONEY, 8: MONEY, 9: MONEY, 10: UNIT})
    return summary, detail


def report_csv(report) -> bytes:
    _, detail = report_sheets(report)
    return _csv([detail.head, *detail.rows])


def report_xlsx(report) -> bytes:
    return _xlsx(list(report_sheets(report)))


def _csv(rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    for row in rows:
        w.writerow(csvsafe.row(row))
    return ("﻿" + buf.getvalue()).encode("utf-8")   # BOM: Excel opens UTF-8 names correctly


def _xlsx(sheets: list[Sheet]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(sheet.name[:31])
        r = 1
        if sheet.title:
            _text(ws.cell(row=1, column=1), sheet.title).font = Font(bold=True, size=12)
            r = 3
        for c, name in enumerate(sheet.head, start=1):
            _text(ws.cell(row=r, column=c), name).font = Font(bold=True)
        ws.freeze_panes = ws.cell(row=r + 1, column=1)
        for row in sheet.rows:
            r += 1
            for c, value in enumerate(row):
                cell = ws.cell(row=r, column=c + 1)
                fmt = sheet.formats.get(c)
                if fmt and value not in (None, ""):
                    try:
                        cell.value = float(Decimal(str(value)))
                        cell.number_format = fmt
                        continue
                    except Exception:
                        pass
                _text(cell, value)
        for c in range(1, len(sheet.head) + 1):
            width = max([len(str(sheet.head[c - 1]))] + [len(str(row[c - 1])) for row in sheet.rows if len(row) >= c])
            ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 9), 48)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _text(cell, value):
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"   # a value starting with "=" stays text, never a formula
    return cell
