"""Month-end close (apps/close): accrual selection by date, estimate methods and confidence, partial
invoicing, adjustments, period locking and versions, exports; vendor statement reading (CSV, XLSX and PDF by
rules, AI with grounding), every reconciliation bucket, resolutions, payments (manual and QuickBooks),
permissions, the statement-sent-as-invoice guard and the synthetic statement generator."""
from __future__ import annotations

import csv
import io
import itertools
import json
import uuid
from datetime import date, datetime, timedelta
from datetime import timezone as dt_timezone
from decimal import Decimal

import httpx
import pytest
from django.urls import reverse
from django.utils import timezone

from apps.accounting.models import PostedBill, QBOConnection, VendorMapping
from apps.close import groups
from apps.close.models import (
    AccrualAdjustment,
    AccrualSnapshot,
    CloseSettings,
    LockedVersionError,
    StatementLine,
    VendorPayment,
    VendorStatement,
)
from apps.close.rules import looks_like_statement
from apps.close.services import accruals, estimates, reconcile, statement_reader
from apps.close.services import statements as statement_service
from apps.core.models import AuditEvent
from apps.documents.models import Document, ExtractedField
from apps.documents.services.ingest import ingest_bytes
from apps.rates.models import Quote, QuoteCharge
from apps.shipments.models import MatchLink, Shipment, ValidationIssue
from apps.shipments.services.containers import make_container
from apps.shipments.services.validation import validate_shipment
from synthetic import month_end
from tests.conftest import PASSWORD

D = Decimal
HL = "Harborlink Logistics LLC"
MD = "Metro Drayage Co."
P = date(2026, 9, 30)
_seq = itertools.count(1)


# --------------------------------------------------------------------------- helpers


def _doc(org, shipment, doc_type, fields, text="", received=None):
    n = next(_seq)
    d = Document.objects.create(organization=org, original_filename=f"{doc_type}-{n}.pdf", sha256=uuid.uuid4().hex,
                                doc_type=doc_type, text=text, status=Document.Status.MATCHED)
    if received:
        Document.objects.filter(pk=d.pk).update(received_at=received)
        d.refresh_from_db()
    for k, v in fields.items():
        ExtractedField.objects.create(document=d, name=k, value=v, confidence=0.99, grounded=True)
    if shipment is not None:
        MatchLink.objects.create(document=d, shipment=shipment, method=MatchLink.Method.EXACT_BL, score=1.0)
    return d


def shipment(org, *, ship="2026-09-10", boxes=2, pol="Ningbo", pod="Long Beach, CA", equipment="40HC",
             with_bl=True, bl_text_extra="", status=Shipment.Status.NEEDS_REVIEW):
    n = next(_seq)
    containers = [make_container("OSL", 200000 + n * 10 + i) for i in range(boxes)]
    bl = f"OSLN{8000000000 + n}"
    s = Shipment.objects.create(organization=org, bl_number=bl if with_bl else "", container_numbers=containers,
                                status=status)
    if with_bl:
        table = "\n".join(f"{c} SL{n}{i} {equipment} 12,000" for i, c in enumerate(containers))
        _doc(org, s, "bill_of_lading", {"carrier_name": "Oceanic Star Line", "bl_number": bl, "issue_date": ship,
                                        "port_of_loading": pol, "port_of_discharge": pod,
                                        "container_numbers": containers},
             text=f"BILL OF LADING\n{table}\n{bl_text_extra}")
    return s


def invoice(org, s, *, vendor=HL, lines=(("Ocean Freight", "4800.00"),), day="2026-09-20", currency="USD",
            number=None, doc_type="freight_invoice"):
    n = next(_seq)
    items = [{"description": desc, "amount": amount} for desc, amount in lines]
    total = sum((D(a) for _, a in lines), D("0.00"))
    return _doc(org, s, doc_type, {
        "vendor_name": vendor, "invoice_number": number or f"INV-{n}", "invoice_date": day, "currency": currency,
        "bl_number": s.bl_number if s else "", "container_numbers": s.container_numbers if s else [],
        "line_items": items, "total_amount": f"{total:.2f}"})


def post(doc, s, day=None):
    return PostedBill.objects.create(organization=doc.organization, document=doc, shipment=s,
                                     request_id=f"sm-{doc.pk}-{uuid.uuid4().hex[:6]}", status=PostedBill.Status.POSTED,
                                     qbo_bill_id=f"Q{doc.pk}", posted_at=day or timezone.now())


def quote(org, vendor=HL, lines=(("thc_destination", "350.00", "container"), ("customs_clearance", "195.00",
                                                                                "shipment")),
          origin="Ningbo", destination="Long Beach, CA", equipment="40HC", currency="USD"):
    q = Quote.objects.create(organization=org, vendor_name=vendor, origin=origin, destination=destination,
                             equipment=equipment, valid_from=date(2026, 1, 1), valid_to=date(2026, 12, 31),
                             currency=currency, reference=f"Q-{next(_seq)}")
    for code, amount, basis in lines:
        QuoteCharge.objects.create(quote=q, code=code, amount=D(amount), basis=basis)
    return q


def settings_for(org, **kw):
    cfg = CloseSettings.for_org(org)
    for k, v in kw.items():
        setattr(cfg, k, v)
    cfg.save()
    return cfg


def lines_for(report, shipment_id=None, kind=None, group=None):
    return [ln for ln in report["lines"] if (shipment_id is None or ln["shipment_id"] == shipment_id)
            and (kind is None or ln["kind"] == kind) and (group is None or ln["group"] == group)]


def login(client, user):
    assert client.login(username=user.username, password=PASSWORD)
    return client


# --------------------------------------------------------------------------- accruals: which lines


def test_selection_by_period_end_date(org):
    settings_for(org, expect_destination="never")
    shipped = shipment(org, ship="2026-09-20")                 # shipped in the period, nothing invoiced
    later = shipment(org, ship="2026-10-02")                   # ships after the period end
    received = shipment(org, ship="2026-09-05")
    in_period = invoice(org, received, day="2026-09-25")       # received, dated in the period, not posted
    late_for_old = invoice(org, shipped, day="2026-10-05", lines=(("Customs Clearance", "195.00"),))
    late_for_new = invoice(org, later, day="2026-10-06")

    report = accruals.build(org, P)
    assert {ln["shipment_id"] for ln in lines_for(report, kind="estimate")} == {shipped.pk}
    received_docs = {ln["document_id"] for ln in lines_for(report, kind="received")}
    # A September shipment's October invoice belongs to September; October's shipment and invoice don't.
    assert received_docs == {in_period.pk, late_for_old.pk}
    assert late_for_new.pk not in received_docs
    late = next(ln for ln in report["lines"] if ln["document_id"] == late_for_old.pk)
    assert "after the period end" in late["basis"]

    october = accruals.build(org, date(2026, 10, 31))
    assert later.pk in {ln["shipment_id"] for ln in october["lines"]}
    assert late_for_new.pk in {ln["document_id"] for ln in lines_for(october, kind="received")}


def test_posted_bills_are_booked_by_bill_date(org):
    settings_for(org, expect_freight="never", expect_destination="never")
    s = shipment(org, ship="2026-09-02")
    booked = invoice(org, s, day="2026-09-15")
    post(booked, s)
    after = invoice(org, s, day="2026-10-03", lines=(("Customs Clearance", "195.00"),))
    post(after, s)
    report = accruals.build(org, P)
    docs = {ln["document_id"]: ln for ln in lines_for(report, kind="received")}
    assert booked.pk not in docs                       # in QuickBooks with a September date: already in the books
    assert docs[after.pk]["status"].startswith("Posted to QuickBooks with bill date 3 Oct 2026")


def test_on_board_date_on_the_bill_of_lading_wins_over_issue_date(org):
    s = shipment(org, ship="2026-10-01", bl_text_extra="Shipped on board: 29 Sep 2026")
    report = accruals.build(org, P)
    assert {ln["shipment_id"] for ln in report["lines"]} == {s.pk}
    assert report["lines"][0]["ship_date"] == "2026-09-29"
    from apps.close.services.history import on_board_date

    assert on_board_date("SHIPPED ON BOARD\nDate: 2026-09-28") == date(2026, 9, 28)
    assert on_board_date("Shipped on board in apparent good order\nPrinted 2026-09-01") is None


def test_credit_notes_and_goods_invoices(org):
    settings_for(org, expect_freight="never", expect_destination="never")
    s = shipment(org, ship="2026-09-05")
    _doc(org, s, "credit_note", {"vendor_name": HL, "credit_note_number": "CN-1", "invoice_date": "2026-09-20",
                                 "currency": "USD", "total_amount": "120.00"})
    goods = invoice(org, s, vendor="Brightway Electronics Co., Ltd.", doc_type="commercial_invoice",
                    lines=(("Speakers", "9000.00"),), day="2026-09-01")
    report = accruals.build(org, P)
    by_doc = {ln["document_id"]: ln for ln in report["lines"]}
    credit = next(ln for ln in report["lines"] if ln["doc_type"] == "Credit note")
    assert credit["amount_home"] == "-120.00" and credit["method_label"] == "Credit note received"
    assert by_doc[goods.pk]["group"] == groups.GOODS and by_doc[goods.pk]["account_name"] == "Inventory"
    settings_for(org, include_goods=False)
    assert goods.pk not in {ln["document_id"] for ln in accruals.build(org, P)["lines"]}


# --------------------------------------------------------------------------- accruals: estimates


def _history(org, n=3, vendor=HL, pol="Ningbo", pod="Long Beach, CA", per_box="300.00", equipment="40HC",
             group_line="Terminal Handling Charge (THC)"):
    for i in range(n):
        h = shipment(org, ship=f"2026-0{3 + i}-10", pol=pol, pod=pod, equipment=equipment)
        invoice(org, h, vendor=vendor, day=f"2026-0{3 + i}-25",
                lines=(("Ocean Freight", "4600.00"), (group_line, f"{D(per_box) * 2 + i * 20:.2f}")))
        h.status = Shipment.Status.POSTED
        h.save()
        for d in h.documents:
            if d.doc_type == "freight_invoice":
                post(d, h, day=datetime(2026, 4, 1, tzinfo=dt_timezone.utc))


def test_estimate_order_quote_then_vendor_median_then_org_median(org):
    settings_for(org, expect_freight="never")
    _history(org)  # Harborlink billed THC on Ningbo to Long Beach 40HC three times: 600, 620, 640 for 2 boxes
    target = shipment(org, ship="2026-09-12")
    invoice(org, target, lines=(("Ocean Freight", "4800.00"),))   # destination charges still to come
    q = quote(org)

    line = lines_for(accruals.build(org, P), target.pk, kind="estimate")[0]
    assert line["method"] == estimates.QUOTE and line["group"] == groups.DESTINATION
    assert line["amount_home"] == "895.00"         # 350 x 2 containers + 195 per shipment
    assert line["confidence_label"] == "High" and q.reference in line["basis"]
    assert line["status"].startswith("Partly invoiced")

    q.archived = True
    q.save()
    line = lines_for(accruals.build(org, P), target.pk, kind="estimate")[0]
    assert line["method"] == estimates.VENDOR_MEDIAN
    assert line["amount_home"] == "620.00"         # median 310.00 per container x 2
    assert line["confidence_label"] == "Medium" and "Median of 3 past shipments billed by Harborlink" in line["basis"]

    other_lane = shipment(org, ship="2026-09-14", pol="Istanbul (Ambarli)")
    invoice(org, other_lane, lines=(("Ocean Freight", "4800.00"),))
    line = lines_for(accruals.build(org, P), other_lane.pk, kind="estimate")[0]
    assert line["method"] == estimates.ORG_MEDIAN and line["confidence_label"] == "Low"
    assert line["vendor_name"] == HL

    settings_for(org, min_history=10)
    line = lines_for(accruals.build(org, P), other_lane.pk, kind="estimate")[0]
    assert line["method"] == estimates.NONE and line["amount_home"] is None
    assert any("no amount yet" in w for w in accruals.build(org, P)["warnings"])


def test_quote_confidence_drops_with_assumptions(org):
    settings_for(org, expect_freight="never")
    quote(org)
    unknown_boxes = shipment(org, ship="2026-09-12", boxes=0)
    line = lines_for(accruals.build(org, P), unknown_boxes.pk, kind="estimate")[0]
    assert line["method"] == estimates.QUOTE and line["confidence"] < 0.9
    assert "1 assumed" in line["basis"]
    assert line["amount_home"] == "545.00"        # 350 x 1 assumed container + 195


def test_partial_invoicing_and_expected_groups(org):
    settings_for(org, expect_delivery="usual")
    for i in range(3):  # most shipments to Long Beach get a trucking invoice
        h = shipment(org, ship=f"2026-0{4 + i}-01")
        invoice(org, h, lines=(("Ocean Freight", "4700.00"), ("Terminal Handling Charge (THC)", "640.00")))
        invoice(org, h, vendor=MD, lines=(("Drayage Port to Warehouse", "1100.00"), ("Chassis Rental", "110.00")))
        h.status = Shipment.Status.POSTED
        h.save()
        for d in h.documents:
            if d.doc_type == "freight_invoice":
                post(d, h, day=datetime(2026, 7, 1, tzinfo=dt_timezone.utc))
    s = shipment(org, ship="2026-09-15")
    invoice(org, s, lines=(("Ocean Freight", "4700.00"), ("Terminal Handling Charge (THC)", "650.00")))
    est = lines_for(accruals.build(org, P), s.pk, kind="estimate")
    assert [ln["group"] for ln in est] == [groups.DELIVERY]      # freight and destination are billed
    assert est[0]["vendor_name"] == MD and est[0]["method"] == estimates.VENDOR_MEDIAN
    assert est[0]["amount_home"] == "1210.00"

    settings_for(org, expect_delivery="never")
    assert not lines_for(accruals.build(org, P), s.pk, kind="estimate")


def test_left_out_rejected_duplicates_undated_and_old(org):
    rejected = shipment(org, ship="2026-09-01", status=Shipment.Status.REJECTED)
    invoice(org, rejected, day="2026-09-10")
    dup_ship = shipment(org, ship="2026-09-01")
    first = invoice(org, dup_ship, number="HA-1", day="2026-09-10")
    copy = invoice(org, dup_ship, number="HA-1", day="2026-09-10")
    ValidationIssue.objects.create(organization=org, shipment=dup_ship, document=copy, code="duplicate_invoice",
                                   severity="error", message="dup", fingerprint="x")
    undated = shipment(org, with_bl=False)
    invoice(org, undated, lines=(("Ocean Freight", "100.00"),), day="2026-09-03")
    old = shipment(org, ship="2025-12-01")
    report = accruals.build(org, P)
    reasons = {(x["shipment_id"], x["reason"].split(",")[0]) for x in report["left_out"]}
    assert (rejected.pk, "Shipment rejected") in reasons
    assert (dup_ship.pk, "Possible duplicate invoice") in reasons
    assert any(sid == undated.pk and r == "No shipping date" for sid, r in reasons)
    assert any(sid == old.pk and r.startswith("Shipped 1 Dec 2025") for sid, r in reasons)
    docs = {ln["document_id"] for ln in report["lines"]}
    assert first.pk in docs and copy.pk not in docs
    assert rejected.pk not in {ln["shipment_id"] for ln in report["lines"]}
    # the undated shipment's own invoice is still accrued by its date
    assert undated.pk in {ln["shipment_id"] for ln in lines_for(report, kind="received")}


def test_accounts_from_vendor_mappings_and_missing_exchange_rates(org):
    settings_for(org, expect_freight="never", expect_destination="never")
    VendorMapping.objects.create(organization=org, vendor_key="harborlink logistics", display_name=HL,
                                 expense_account_id="88", expense_account_name="Freight In")
    s = shipment(org, ship="2026-09-01")
    invoice(org, s, day="2026-09-10", lines=(("Ocean Freight", "1000.00"),))
    invoice(org, s, vendor="Swift Cargo Forwarding Inc.", day="2026-09-11", lines=(("Ocean Freight", "500.00"),))
    eur = invoice(org, s, vendor="Anatolia Textile Export A.S.", doc_type="commercial_invoice", currency="EUR",
                  day="2026-09-01", lines=(("Towels", "2000.00"),))
    report = accruals.build(org, P)
    accounts = {a["account_name"]: a for a in report["by_account"]}
    assert accounts["Freight In"]["total"] == "1000.00" and accounts["Freight In"]["account_id"] == "88"
    assert accounts["Freight expense"]["total"] == "500.00"
    assert report["totals"]["total"] == "1500.00"
    assert any("EUR" in w for w in report["warnings"])
    assert next(ln for ln in report["lines"] if ln["document_id"] == eur.pk)["amount_home"] is None
    org.fx_rates = {"EUR": "1.10"}
    org.save()
    assert accruals.build(org, P)["totals"]["total"] == "3700.00"


# --------------------------------------------------------------------------- adjustments


def test_adjustments_exclude_or_set_amount(client, org, approver, user):
    settings_for(org, expect_freight="never")
    s = shipment(org, ship="2026-09-12")
    login(client, approver)
    url = reverse("close:adjust")
    r = client.post(url, {"shipment": s.pk, "group": "destination", "action": "amount", "amount": "1,850.00",
                          "vendor": HL, "note": "Quoted by email", "period": "2026-09-30"})
    assert r.status_code == 302
    line = lines_for(accruals.build(org, P), s.pk)[0]
    assert line["method"] == estimates.MANUAL and line["amount_home"] == "1850.00"
    assert line["confidence_label"] == "Set by a person" and "Quoted by email" in line["basis"]

    client.post(url, {"shipment": s.pk, "group": "destination", "action": "exclude", "note": "Customer collects",
                      "period": "2026-09-30"})
    report = accruals.build(org, P)
    assert not lines_for(report, s.pk)
    assert any("not needed" in x["reason"] and "Customer collects" in x["detail"] for x in report["left_out"])
    adj = AccrualAdjustment.objects.get(shipment=s)
    assert AuditEvent.objects.filter(action="close.accrual_adjusted", object_id=str(adj.pk)).count() == 2

    # a note is required, and reviewers can't adjust
    r = client.post(url, {"shipment": s.pk, "group": "destination", "action": "exclude", "note": "",
                          "period": "2026-09-30"}, follow=True)
    assert "Add a note" in r.content.decode()
    client.logout()
    login(client, user)
    assert client.post(url, {"shipment": s.pk, "group": "destination", "action": "exclude", "note": "x"}
                       ).status_code == 403
    client.logout()
    login(client, approver)
    client.post(reverse("close:remove_adjustment", args=[adj.pk]), {"period": "2026-09-30"})
    assert not AccrualAdjustment.objects.exists()
    assert lines_for(accruals.build(org, P), s.pk)[0]["method"] != estimates.MANUAL


# --------------------------------------------------------------------------- locking and versions


def test_lock_period_versions_are_read_only(client, org, approver, user):
    settings_for(org, expect_destination="never")
    s = shipment(org, ship="2026-09-12")
    invoice(org, s, day="2026-09-20")
    login(client, approver)
    r = client.post(reverse("close:lock"), {"period": "2026-09-30", "expected_version": "0", "note": ""})
    assert r.status_code == 302
    v1 = AccrualSnapshot.objects.get(organization=org, period_end=P)
    assert v1.version == 1 and v1.total == D("4800.00") and v1.intact
    assert AuditEvent.objects.filter(action="close.period_locked", object_id=str(v1.pk)).exists()

    with pytest.raises(LockedVersionError):
        v1.note = "changed"
        v1.save()
    with pytest.raises(LockedVersionError):
        v1.delete()

    # The locked page shows the stored report even after new invoices arrive.
    invoice(org, s, day="2026-10-02", lines=(("Customs Clearance", "195.00"),))
    page = client.get(reverse("close:accruals"), {"period": "2026-09-30"}).content.decode()
    assert "Locked: version 1" in page and "4,995.00" not in page
    live = client.get(reverse("close:accruals"), {"period": "2026-09-30", "live": "1"}).content.decode()
    assert "Live preview" in live and "1 line appeared" in live

    # A second lock needs a reason, and a stale page can't stack a version.
    r = client.post(reverse("close:lock"), {"period": "2026-09-30", "expected_version": "1", "note": ""}, follow=True)
    assert "Say why the booked amount changes" in r.content.decode()
    r = client.post(reverse("close:lock"), {"period": "2026-09-30", "expected_version": "0", "note": "late invoice"},
                    follow=True)
    assert "locked while you were looking" in r.content.decode()
    client.post(reverse("close:lock"), {"period": "2026-09-30", "expected_version": "1", "note": "Late customs bill"})
    v2 = AccrualSnapshot.objects.get(organization=org, period_end=P, version=2)
    assert v2.total == D("4995.00") and v2.note == "Late customs bill"
    v1.refresh_from_db()
    assert v1.total == D("4800.00") and v1.intact

    # Reviewers can't lock; a period that hasn't ended can't be locked.
    client.logout()
    login(client, user)
    assert client.post(reverse("close:lock"), {"period": "2026-09-30"}).status_code == 403
    client.logout()
    login(client, approver)
    future = (timezone.localdate() + timedelta(days=40)).isoformat()
    r = client.post(reverse("close:lock"), {"period": future, "expected_version": "0"}, follow=True)
    assert "ended yet. Lock a period after its last day" in r.content.decode()


# --------------------------------------------------------------------------- exports


def test_csv_and_journal_exports_from_locked_version(client, org, approver):
    from openpyxl import load_workbook

    settings_for(org, expect_destination="never")
    s = shipment(org, ship="2026-09-12")
    invoice(org, s, vendor="=HYPERLINK(\"http://x\")", day="2026-09-20", lines=(("Ocean Freight", "1000.00"),))
    unbilled = shipment(org, ship="2026-09-14")
    quote(org, lines=(("ocean_freight", "2000.00", "container"),))
    login(client, approver)
    client.post(reverse("close:lock"), {"period": "2026-09-30", "expected_version": "0"})
    invoice(org, unbilled, day="2026-10-01")  # arrives after locking: version 1 must not change

    r = client.get(reverse("close:accruals_csv"), {"period": "2026-09-30"})
    assert r.status_code == 200 and "accruals-2026-09-30-v1.csv" in r["Content-Disposition"]
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][:3] == ["Period end", "Version", "Type"]
    vendors = {row[6] for row in rows[1:] if len(row) > 6}
    assert "'=HYPERLINK(\"http://x\")" in vendors        # formula guard
    assert rows[-1][13] == "5000.00"                      # 1000 received + 2 x 2000 quoted

    r = client.get(reverse("close:accruals_journal"), {"period": "2026-09-30"})
    wb = load_workbook(io.BytesIO(r.content))
    je, rev, lines = wb["Journal entry"], wb["Reversing entry"], wb["Lines"]
    rows = list(je.iter_rows(min_row=2, values_only=True))
    debits = sum(D(str(r[4])) for r in rows if r[4] is not None)
    credits = sum(D(str(r[5])) for r in rows if r[5] is not None)
    assert debits == credits == D("5000")
    assert rows[-1][2] == "Accrued liabilities" and rows[-1][0] == "ACCR-2026-09-v1"
    rrows = list(rev.iter_rows(min_row=2, values_only=True))
    assert rrows[-1][4] == 5000 and rrows[-1][5] is None          # the reversal debits the liability
    assert rrows[0][1].date() == date(2026, 10, 1) and rrows[0][0] == "ACCR-2026-09-v1-R"
    assert lines.max_row == 3                                      # header + 2 lines, as locked
    notes = {r[0]: r[1] for r in wb["Notes"].iter_rows(values_only=True)}
    assert notes["Version"] == "Version 1, locked" and len(notes["Report fingerprint (SHA-256)"]) == 64
    assert AuditEvent.objects.filter(action="close.accruals_exported").count() == 2

    live = client.get(reverse("close:accruals_csv"), {"period": "2026-09-30", "live": "1"})
    assert "live.csv" in live["Content-Disposition"]
    assert client.get(reverse("close:accruals_csv"), {"period": "2026-09-30", "version": "7"}).status_code == 404


# --------------------------------------------------------------------------- permissions and pages


def test_month_end_permissions_and_navigation(client, org, viewer, user, approver, admin_user):
    for who in (viewer, user):
        login(client, who)
        for name in ("close:accruals", "close:statements"):
            assert client.get(reverse(name)).status_code == 403
        assert "Month-end" not in client.get(reverse("core:dashboard")).content.decode()
        client.logout()
    login(client, approver)
    page = client.get(reverse("close:accruals")).content.decode()
    assert "Not locked yet" in page and "Lock version 1" in page
    dash = client.get(reverse("core:dashboard")).content.decode()
    assert "Month-end" in dash and "Not locked" in dash
    assert client.get(reverse("close:statements")).status_code == 200
    assert client.get(reverse("close:settings")).status_code == 403
    client.logout()
    login(client, admin_user)
    r = client.post(reverse("close:settings"), {
        "accrued_account": "Accrued freight", "accrued_account_id": "2100", "freight_account": "Freight in",
        "goods_account": "Inventory", "include_goods": "on", "expect_freight": "always",
        "expect_destination": "usual", "expect_delivery": "never", "lookback_days": "120", "min_history": "2"})
    assert r.status_code == 302
    cfg = CloseSettings.objects.get(organization=org)
    assert cfg.accrued_account == "Accrued freight" and cfg.lookback_days == 120
    assert AuditEvent.objects.filter(action="close.settings_updated").exists()
    r = client.post(reverse("close:settings"), {"accrued_account": "", "freight_account": "x", "goods_account": "y",
                                                "expect_freight": "always", "expect_destination": "always",
                                                "expect_delivery": "usual", "lookback_days": "2", "min_history": "0"})
    assert r.status_code == 200 and "Use a window between 7 and 1,095 days." in r.content.decode()


def test_other_organizations_are_invisible(client, org, approver):
    from apps.core.models import Organization

    other = Organization.objects.create(name="Other", slug="other")
    st = VendorStatement.objects.create(organization=other, vendor_name=HL, vendor_key="harborlink logistics",
                                        sha256="x" * 64, original_filename="s.csv", source_format="csv")
    login(client, approver)
    assert client.get(reverse("close:statement", args=[st.pk])).status_code == 404
    assert client.post(reverse("close:statement_rematch", args=[st.pk])).status_code == 404


# --------------------------------------------------------------------------- statements: reading


@pytest.fixture
def harborlink_scenario(dataset):
    return month_end.scenario(month_end.invoices_from_ground_truth(dataset))


@pytest.mark.parametrize("fmt", ["pdf", "xlsx", "csv"])
def test_statement_read_by_rules(harborlink_scenario, fmt):
    s = harborlink_scenario
    data = {"pdf": month_end.statement_pdf, "xlsx": month_end.statement_xlsx, "csv": month_end.statement_csv}[fmt](s)
    p = statement_reader.read(f"statement.{fmt}", data)
    assert p.reader == "rules" and p.vendor_name == HL and p.currency == "USD"
    assert p.statement_date == s["statement_date"] and p.closing_balance == s["closing_balance"]
    assert [ln.number for ln in p.lines] == [ln.number for ln in s["lines"]]
    assert [ln.amount for ln in p.lines] == [ln.amount for ln in s["lines"]]
    assert all(ln.kind == "invoice" for ln in p.lines)
    assert p.lines[0].reference == s["lines"][0].reference


def test_statement_reader_types_signs_and_open_amounts():
    text = ("Pacific Crest Lines\nSTATEMENT OF ACCOUNT\nStatement date: 30 Sep 2026\n"
            "Date Type Number Reference Debit Credit Balance\n"
            "01 Sep 2026 Balance brought forward 1,000.00\n"
            "03 Sep 2026 Invoice PC-1001 OSLN1 500.00 1,500.00\n"
            "10 Sep 2026 Credit note CN-55 PC-1001 (50.00) 1,450.00\n"
            "15 Sep 2026 Payment received WIRE-9 1,000.00 450.00\n"
            "Balance due: 450.00\nCurrent: 450.00 30 days: 0.00\n")
    p = statement_reader.read_pdf_text(text)
    statement_reader._finish(p)
    kinds = [(ln.kind, ln.number, ln.amount) for ln in p.lines]
    assert kinds == [("opening", "", D("1000.00")), ("invoice", "PC-1001", D("500.00")),
                     ("credit", "CN-55", D("-50.00")), ("payment", "WIRE-9", D("-1000.00"))]
    assert p.opening_balance == D("1000.00") and p.closing_balance == D("450.00")
    assert p.statement_date == date(2026, 9, 30) and p.vendor_name == "Pacific Crest Lines"

    # open-items sheet: the balance column is what is still open on each line
    rows = [["Invoice", "Date", "Amount", "Balance"], ["A-1", "2026-09-01", "100.00", "40.00"],
            ["A-2", "2026-09-05", "200.00", "200.00"], ["A-3", "2026-09-07", "300.00", "10.00"]]
    from apps.intake.services.sheets import Sheet

    p = statement_reader.read_sheet_rows([Sheet("s", rows)])
    statement_reader._finish(p)
    assert [ln.amount for ln in p.lines] == [D("40.00"), D("200.00"), D("10.00")]
    assert any("still open" in n for n in p.notes)


def test_statement_rejects_unsupported_and_unreadable_files(client, org, approver):
    from synthetic.extra import photo

    login(client, approver)
    pdf = month_end.statement_pdf(month_end.scenario([
        {"invoice_number": "A-1", "invoice_date": "2026-09-01", "bl_number": "", "total_amount": "10.00"},
        {"invoice_number": "A-2", "invoice_date": "2026-09-02", "bl_number": "", "total_amount": "20.00"}]))
    png = photo(pdf)
    r = client.post(reverse("close:statement_upload"),
                    {"file": _named(png, "scan.png")}, follow=True)
    assert "statements are read from PDF, Excel (.xlsx) or CSV files" in r.content.decode()
    r = client.post(reverse("close:statement_upload"), {"file": _named(b"\x00\x01binary", "statement.txt")}, follow=True)
    assert "this isn&#x27;t a file ShipMatch can read" in r.content.decode()
    r = client.post(reverse("close:statement_upload"), {"file": _named(b"Just a note\nno amounts here\n", "x.csv")},
                    follow=True)
    assert "No statement lines could be read" in r.content.decode()
    assert not VendorStatement.objects.exists()


def _named(content: bytes, name: str):
    from django.core.files.uploadedfile import SimpleUploadedFile

    return SimpleUploadedFile(name, content)


def test_statement_read_by_ai_is_grounded_with_rules_fallback(settings, monkeypatch, harborlink_scenario):
    from apps.documents.services import llm

    s = harborlink_scenario
    data = month_end.statement_csv(s)
    settings.EXTRACTION_PROVIDER = "anthropic"
    first = s["lines"][0]
    answer = {"vendor_name": HL, "statement_date": "2026-07-18", "currency": "usd", "opening_balance": None,
              "closing_balance": float(s["closing_balance"]), "lines": [
                  {"type": "invoice", "invoice_number": first.number, "date": first.day.isoformat(),
                   "reference": first.reference, "amount": float(first.amount), "balance": None},
                  {"type": "invoice", "invoice_number": "INVENTED-1", "date": "2026-07-01", "reference": None,
                   "amount": 99999.99, "balance": None}]}
    calls = []

    def fake(system, user, schema, **kw):
        calls.append((schema, kw))
        return answer

    monkeypatch.setattr(llm, "structured_call", fake)
    p = statement_reader.read("statement.csv", data)
    assert p.reader == "anthropic" and [ln.number for ln in p.lines] == [first.number]
    assert any("not printed in the statement" in n for n in p.notes)
    schema = calls[0][0]
    assert schema["additionalProperties"] is False and schema["properties"]["lines"]["items"][
        "additionalProperties"] is False
    assert "$ref" not in json.dumps(schema) and "pattern" not in json.dumps(schema)

    def broken(*a, **kw):
        raise llm.LLMError("down")

    monkeypatch.setattr(llm, "structured_call", broken)
    p = statement_reader.read("statement.csv", data)
    assert p.reader == "rules" and len(p.lines) == len(s["lines"])
    assert any("could not read this statement" in n for n in p.notes)


# --------------------------------------------------------------------------- statements: the synthetic scenario


@pytest.mark.parametrize("fmt", ["pdf", "xlsx"])
def test_synthetic_statement_reconciles_with_planted_differences(client, org, approver, dataset, fmt):
    for f in sorted(p.name for p in (dataset / "pdf").iterdir()):
        ingest_bytes(org, f, (dataset / "pdf" / f).read_bytes(), process="sync")
    s = month_end.scenario(month_end.invoices_from_ground_truth(dataset))
    credit_pdf, _ = s["credit_note"]
    ingest_bytes(org, "credit-note.pdf", credit_pdf, process="sync")
    data = month_end.statement_pdf(s) if fmt == "pdf" else month_end.statement_xlsx(s)
    documents_before = Document.objects.count()

    login(client, approver)
    r = client.post(reverse("close:statement_upload"), {"file": _named(data, f"harborlink.{fmt}")})
    st = VendorStatement.objects.get(organization=org)
    assert r.status_code == 302 and r["Location"] == reverse("close:statement", args=[st.pk])
    assert st.vendor_name == HL and st.status == VendorStatement.Status.READY
    buckets = {}
    for item in st.items.select_related("line"):
        buckets.setdefault(item.bucket, set()).add(item.line.number if item.line else item.label.split()[-1])
    exp = s["expected"]
    # The dataset has one scanned invoice ShipMatch can't read without OCR: the statement shows it as never received.
    truth = json.loads((dataset / "ground_truth.json").read_text())
    scanned = {d["fields"].get("invoice_number") for d in truth["documents"] if d["scanned"]} & set(exp["matched"])
    assert buckets["missing"] == set(exp["missing"]) | scanned
    unread = st.items.get(bucket="missing", line__number__in=scanned) if scanned else None
    assert unread is None or "couldn't read yet" in unread.explanation
    assert buckets["amount_differs"] == set(exp["amount_differs"])
    assert buckets["credit_not_applied"] == set(exp["credit_not_applied"])
    assert buckets["matched"] == set(exp["matched"]) - scanned
    assert set(buckets) == {"matched", "missing", "amount_differs", "credit_not_applied"}
    unread_total = sum((ln.amount for ln in s["lines"] if ln.number in scanned), D("0.00"))
    assert D(st.summary["difference"]) == D(exp["difference"]) + unread_total
    assert st.summary["unexplained"] == "0.00"
    # A statement is never a document: nothing was classified, matched or queued for posting.
    assert Document.objects.count() == documents_before

    page = client.get(reverse("close:statement", args=[st.pk])).content.decode()
    assert "Request a copy" in page and "mailto:" in page and "Why the balances differ" in page

    # Uploading the same file again shows the same statement.
    r = client.post(reverse("close:statement_upload"), {"file": _named(data, "again.pdf")})
    assert VendorStatement.objects.count() == 1


# --------------------------------------------------------------------------- statements: every bucket


def _statement(org, lines, *, vendor=HL, day=date(2026, 9, 30), closing=None, opening=None):
    st = VendorStatement.objects.create(organization=org, vendor_name=vendor, vendor_key=reconcile.vendor_key(vendor),
                                        statement_date=day, currency="USD", sha256=uuid.uuid4().hex,
                                        original_filename="s.csv", source_format="csv", closing_balance=closing,
                                        opening_balance=opening)
    for n, (kind, number, when, ref, amount) in enumerate(lines, 1):
        StatementLine.objects.create(statement=st, position=n, kind=kind, number=number, date=when, reference=ref,
                                     amount=D(amount))
    return st


def _buckets(st):
    reconcile.run(st)
    st.refresh_from_db()
    out = {}
    for i in st.items.all():
        out.setdefault(i.bucket, []).append(i)
    return out


def test_every_reconciliation_bucket(org):
    s1 = shipment(org, ship="2026-08-01")
    invoice(org, s1, number="HA-100", day="2026-08-20", lines=(("Ocean Freight", "1000.00"),))
    invoice(org, s1, number="HA-101", day="2026-08-21", lines=(("Ocean Freight", "500.00"),))
    invoice(org, s1, number="HA-102", day="2026-08-22", lines=(("Ocean Freight", "300.00"),))   # not on statement
    _doc(org, s1, "credit_note", {"vendor_name": HL, "credit_note_number": "CN-9", "invoice_date": "2026-08-25",
                                  "currency": "USD", "total_amount": "40.00", "original_invoice_number": "HA-999"})
    rejected = shipment(org, ship="2026-08-01", status=Shipment.Status.REJECTED)
    invoice(org, rejected, number="HA-103", day="2026-08-23", lines=(("Ocean Freight", "250.00"),))
    VendorPayment.objects.create(organization=org, vendor_name=HL, paid_on=date(2026, 9, 28), amount=D("700.00"),
                                 currency="USD", reference="WIRE-1")
    VendorPayment.objects.create(organization=org, vendor_name=HL, paid_on=date(2026, 9, 10), amount=D("90.00"),
                                 currency="USD", reference="WIRE-0")
    st = _statement(org, [
        ("invoice", "HA-100", date(2026, 8, 20), "", "1000.00"),         # matched
        ("invoice", "ha 101", date(2026, 8, 21), "", "525.00"),          # amount differs (+25)
        ("invoice", "HA-100", date(2026, 8, 20), "", "1000.00"),         # duplicate on the statement
        ("invoice", "HA-555", date(2026, 9, 2), "OSLN1", "80.00"),       # never received
        ("invoice", "HA-103", date(2026, 8, 23), "", "250.00"),          # rejected in ShipMatch
        ("payment", "", date(2026, 9, 12), "WIRE-0", "-90.00"),          # matched payment
        ("payment", "", date(2026, 9, 15), "CHK-77", "-60.00"),          # payment ShipMatch doesn't know
    ], closing=D("2700.00"))
    b = _buckets(st)
    assert {i.label for i in b["matched"]} == {"Invoice HA-100", "Payment on line 6"}
    differs = {i.label: i for i in b["amount_differs"]}
    assert differs["Invoice ha 101"].effect == D("25.00")
    assert differs["Invoice HA-103"].shipmatch_amount == D("0.00") and "rejected" in differs["Invoice HA-103"].explanation
    assert [i.line.position for i in b["duplicate"]] == [3]
    assert [i.line.number for i in b["missing"]] == ["HA-555"]
    assert [i.document.field("invoice_number") for i in b["not_on_statement"]] == ["HA-102"]
    assert [i.label for i in b["credit_not_applied"]] == ["Credit note CN-9"]
    assert [i.payment.reference for i in b["payment_not_applied"]] == ["WIRE-1"]
    assert "may still be on its way" in b["payment_not_applied"][0].explanation
    assert [i.statement_amount for i in b["payment_unknown"]] == [D("60.00")]
    assert [i.effect for i in b["arithmetic"]] == [D("2700.00") - D("2705.00")]
    summ = st.summary
    # vendor: 2700 printed; ShipMatch: 1000+500+300+0 - 40 - 700 - 90 = 970
    assert summ["statement_balance"] == "2700.00" and summ["shipmatch_balance"] == "970.00"
    assert summ["difference"] == "1730.00" and summ["unexplained"] == "0.00"


def test_credit_applied_by_vendor_is_matched(org):
    s = shipment(org, ship="2026-08-01")
    invoice(org, s, number="HA-200", day="2026-08-20", lines=(("Ocean Freight", "1000.00"),))
    _doc(org, s, "credit_note", {"vendor_name": HL, "credit_note_number": "CN-1", "invoice_date": "2026-08-25",
                                 "currency": "USD", "total_amount": "150.00", "original_invoice_number": "HA-200"})
    st = _statement(org, [("invoice", "HA-200", date(2026, 8, 20), "", "850.00")])
    b = _buckets(st)
    assert set(b) == {"matched"} and "applied credit note cn-1" in b["matched"][0].explanation.lower()
    assert st.summary["difference"] == "0.00" and st.summary["unexplained"] == "0.00"


def test_open_items_statement_settles_paid_invoices_and_opening_balance(org):
    s = shipment(org, ship="2026-07-01")
    paid = invoice(org, s, number="HA-300", day="2026-07-20", lines=(("Ocean Freight", "640.00"),))
    invoice(org, s, number="HA-301", day="2026-08-20", lines=(("Ocean Freight", "410.00"),))
    VendorPayment.objects.create(organization=org, vendor_name=HL, paid_on=date(2026, 8, 30), amount=D("640.00"),
                                 currency="USD", reference="ACH-5",
                                 allocations=[{"document_id": paid.pk, "invoice_number": "HA-300", "amount": "640.00"}])
    st = _statement(org, [("invoice", "HA-301", date(2026, 8, 20), "", "410.00")])
    b = _buckets(st)
    assert set(b) == {"matched", "settled"} and len(b["settled"]) == 2
    assert st.summary["shape"] == "open_items" and st.summary["difference"] == "0.00"

    # An activity statement that starts from a balance: earlier items are compared as one opening balance.
    st2 = _statement(org, [("opening", "", date(2026, 8, 1), "Opening balance", "700.00"),
                           ("invoice", "HA-301", date(2026, 8, 20), "", "410.00")], opening=D("700.00"))
    b = _buckets(st2)
    assert b["opening"][0].effect == D("60.00")            # vendor 700 vs ShipMatch 640 before 20 Aug
    assert "payment_not_applied" in b and st2.summary["unexplained"] == "0.00"


def test_vendor_chosen_by_person_and_resolutions_survive_rematch(client, org, approver, user, viewer):
    s = shipment(org, ship="2026-08-01")
    invoice(org, s, number="HA-400", day="2026-08-20", lines=(("Ocean Freight", "100.00"),))
    content = (b"Statement date,2026-09-30\n\nInvoice No.,Date,Amount\nHA-400,2026-08-20,100.00\n"
               b"HA-401,2026-09-01,75.00\n")
    login(client, approver)
    client.post(reverse("close:statement_upload"), {"file": _named(content, "statement.csv")})
    st = VendorStatement.objects.get()
    assert st.status == VendorStatement.Status.NEEDS_VENDOR and not st.items.exists()
    assert "Which vendor is this statement from?" in client.get(reverse("close:statement", args=[st.pk])).content.decode()
    client.post(reverse("close:statement_edit", args=[st.pk]), {"vendor": HL, "statement_date": "2026-09-30",
                                                                "currency": "usd", "closing_balance": ""})
    st.refresh_from_db()
    assert st.status == VendorStatement.Status.READY and st.vendor_key == "harborlink logistics"
    missing = st.items.get(bucket="missing")

    for who in (user, viewer):
        client.logout()
        login(client, who)
        assert client.post(reverse("close:item_resolve", args=[missing.pk]), {"note": "x"}).status_code == 403
    client.logout()
    login(client, approver)
    r = client.post(reverse("close:item_resolve", args=[missing.pk]), {"note": ""}, follow=True)
    assert "Add a note" in r.content.decode()
    client.post(reverse("close:item_resolve", args=[missing.pk]), {"note": "Copy requested on 2 Oct"})
    client.post(reverse("close:statement_rematch", args=[st.pk]))
    missing.refresh_from_db()
    assert missing.resolved and missing.resolution_note == "Copy requested on 2 Oct"
    assert AuditEvent.objects.filter(action="close.item_resolved", object_id=str(missing.pk)).exists()
    client.post(reverse("close:item_resolve", args=[missing.pk]), {"reopen": "1"})
    missing.refresh_from_db()
    assert not missing.resolved

    # When the invoice arrives the finding goes away.
    invoice(org, s, number="HA-401", day="2026-09-01", lines=(("Customs Clearance", "75.00"),))
    client.post(reverse("close:statement_rematch", args=[st.pk]))
    assert not st.items.filter(bucket="missing").exists()

    r = client.get(reverse("close:statement_export", args=[st.pk]))
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][2] == "Finding" and any(row and row[0] == "Difference" for row in rows)

    client.post(reverse("close:statement_delete", args=[st.pk]))
    assert not VendorStatement.objects.exists()
    assert AuditEvent.objects.filter(action="close.statement_deleted").exists()


def test_request_copy_mailto_uses_vendor_contact(org):
    from urllib.parse import unquote

    from apps.disputes.models import VendorContact

    VendorContact.objects.create(organization=org, vendor_key="harborlink logistics", vendor_name=HL,
                                 email="billing@harborlink.example")
    st = _statement(org, [("invoice", "HA-777", date(2026, 9, 2), "OSLN55", "80.00")])
    item = _buckets(st)["missing"][0]
    link = unquote(reconcile.copy_request_mailto(org, st, item))
    assert link.startswith("mailto:billing@harborlink.example?subject=Copy of invoice HA-777 requested")
    assert "USD 80.00" in link and "OSLN55" in link and org.name in link


# --------------------------------------------------------------------------- payments


def test_record_payment_links_invoices_and_rematches(client, org, approver):
    s = shipment(org, ship="2026-08-01")
    invoice(org, s, number="HA-500", day="2026-08-20", lines=(("Ocean Freight", "640.00"),))
    st = _statement(org, [("invoice", "HA-500", date(2026, 8, 20), "", "640.00"),
                          ("payment", "", date(2026, 9, 5), "WIRE-77", "-640.00")])
    reconcile.run(st)
    assert st.items.filter(bucket="payment_unknown").exists()
    login(client, approver)
    r = client.post(reverse("close:payment_add"), {"vendor": HL, "paid_on": "2026-09-05", "amount": "640",
                                                   "currency": "usd", "reference": "WIRE-77",
                                                   "invoices": "HA-500, HA-999", "statement": st.pk}, follow=True)
    assert "has no invoice HA-999" in r.content.decode()
    pay = VendorPayment.objects.get()
    assert pay.allocations[0]["invoice_number"] == "HA-500" and pay.allocations[0]["document_id"]
    assert not st.items.filter(bucket="payment_unknown").exists()          # matched automatically
    r = client.post(reverse("close:payment_add"), {"vendor": HL, "paid_on": "2099-01-01", "amount": "1",
                                                   "currency": "USD"}, follow=True)
    assert "in the future" in r.content.decode()
    r = client.post(reverse("close:payment_add"), {"vendor": HL, "paid_on": "2026-09-01", "amount": "-5",
                                                   "currency": "USD"}, follow=True)
    assert "more than zero" in r.content.decode()


def test_payments_read_from_quickbooks(org, approver, settings):
    s = shipment(org, ship="2026-08-01")
    inv = invoice(org, s, number="HA-600", day="2026-08-20", lines=(("Ocean Freight", "640.00"),))
    PostedBill.objects.create(organization=org, document=inv, shipment=s, request_id="r-600",
                              status=PostedBill.Status.POSTED, qbo_bill_id="501")
    VendorMapping.objects.create(organization=org, vendor_key="harborlink logistics", display_name=HL,
                                 qbo_vendor_id="77")
    conn = QBOConnection.objects.create(organization=org, realm_id="123", access_token="tok", refresh_token="ref",
                                        access_expires_at=timezone.now() + timedelta(hours=1), home_currency="USD")
    queries = []

    def handler(request):
        queries.append(request.url.params["query"])
        return httpx.Response(200, json={"QueryResponse": {"BillPayment": [{
            "Id": "9001", "TxnDate": "2026-09-05", "TotalAmt": 640.0, "DocNumber": "WIRE-88",
            "VendorRef": {"value": "77", "name": "Harborlink Logistics"}, "CurrencyRef": {"value": "USD"},
            "Line": [{"Amount": 640.0, "LinkedTxn": [{"TxnId": "501", "TxnType": "Bill"}]}]}]}})

    from apps.accounting.services.quickbooks import QBOClient

    client = QBOClient(conn, http=httpx.Client(transport=httpx.MockTransport(handler)))
    result = statement_service.sync_quickbooks_payments(org, approver, client=client, today=date(2026, 10, 3))
    assert result["created"] == 1 and "from BillPayment where TxnDate >= '2026-04-06'" in queries[0]
    pay = VendorPayment.objects.get()
    assert pay.vendor_name == HL and pay.amount == D("640.00") and pay.source == "quickbooks"
    assert pay.allocations[0]["document_id"] == inv.pk and pay.allocations[0]["invoice_number"] == "HA-600"
    again = statement_service.sync_quickbooks_payments(org, approver, client=client, today=date(2026, 10, 3))
    assert again == {**again, "created": 0, "updated": 1} and VendorPayment.objects.count() == 1

    settings.DEMO_MODE, settings.DEMO_SEND_OUTSIDE = True, False
    with pytest.raises(statement_service.SyncError, match="public demo"):
        statement_service.sync_quickbooks_payments(org, approver, client=client)


# --------------------------------------------------------------------------- the invoice pipeline guard


def test_a_statement_in_the_invoice_pipeline_blocks_approval(org):
    text = ("Harborlink Logistics LLC STATEMENT OF ACCOUNT\nStatement date: 2026-09-30\n"
            "2026-09-01 Invoice HA-1 500.00 500.00\nBalance due: 500.00\n")
    assert looks_like_statement(text)
    assert not looks_like_statement("Harborlink Logistics LLC FREIGHT INVOICE\nInvoice No.: HA-1\nTotal: 5.00")
    assert not looks_like_statement("INVOICE STATEMENT\nbalance due 5.00")
    s = shipment(org, ship="2026-09-01")
    doc = invoice(org, s, number="HA-1", day="2026-09-30")
    Document.objects.filter(pk=doc.pk).update(text=text)
    validate_shipment(s)
    issue = s.issues.get(code="looks_like_statement")
    assert issue.severity == "error" and issue.title == "Looks like a vendor statement"


def test_synthetic_month_end_helpers_are_opt_in(tmp_path):
    from synthetic.generator import generate

    generate(tmp_path / "a", n_shipments=8, seed=5)
    names = sorted(p.name for p in (tmp_path / "a" / "pdf").iterdir())
    assert not any("statement" in n or n.startswith("ME") for n in names)
    files = month_end.arrived_not_invoiced(date(2026, 9, 25))
    assert [f[0] for f in files] == ["ME011_commercial_invoice.pdf", "ME011_bill_of_lading.pdf",
                                     "ME012_commercial_invoice.pdf", "ME012_bill_of_lading.pdf",
                                     "ME012_freight_invoice.pdf"]
    assert files == month_end.arrived_not_invoiced(date(2026, 9, 25))     # same bytes every run


def test_arrived_not_invoiced_shipments_are_estimated(org):
    for name, pdf, _ in month_end.arrived_not_invoiced(date(2026, 9, 25)):
        ingest_bytes(org, name, pdf, process="sync")
    quote(org, lines=(("ocean_freight", "2350.00", "container"), ("thc_destination", "350.00", "container"),
                      ("customs_clearance", "195.00", "shipment")))
    report = accruals.build(org, P)
    est = lines_for(report, kind="estimate")
    assert len({ln["shipment_id"] for ln in est}) == 2
    by_group = sorted((ln["group"], ln["amount_home"]) for ln in est)
    # no freight invoice: both groups from the quote; ocean freight billed: only destination charges
    assert by_group == [("destination", "895.00"), ("destination", "895.00"), ("freight", "4700.00")]


def test_close_demo_command_loads_a_statement_with_planted_differences(org, tmp_path, settings):
    from django.core.management import call_command

    for n, (number, day, amount) in enumerate((("HA-710001", "2026-08-03", "2400.00"),
                                               ("HA-710002", "2026-08-19", "3100.00"))):
        s = shipment(org, ship=f"2026-07-2{n}")
        invoice(org, s, number=number, day=day, lines=(("Ocean Freight", amount),))
    out = tmp_path / "out"
    call_command("close_demo", org=org.slug, out=str(out), no_shipments=True)
    assert {p.name for p in out.iterdir()} == {
        "harborlink-logistics-statement.pdf", "harborlink-logistics-statement.xlsx",
        "harborlink-logistics-statement.csv", "harborlink-logistics-credit-note.pdf"}
    st = VendorStatement.objects.get(organization=org)
    assert st.summary["difference"] == "730.00" and st.summary["unexplained"] == "0.00"
    assert {i.bucket for i in st.items.all()} == {"matched", "missing", "amount_differs", "credit_not_applied"}
    call_command("close_demo", org=org.slug, out=str(out), no_shipments=True)   # safe to repeat
    assert VendorStatement.objects.count() == 1


def test_oversized_statement_upload_is_refused(client, org, approver, settings):
    settings.INTAKE_MAX_FILE_MB = 0
    login(client, approver)
    r = client.post(reverse("close:statement_upload"), {"file": _named(b"Invoice No.,Amount\nA-1,5.00\n", "s.csv")},
                    follow=True)
    assert "larger than 0 MB" in r.content.decode() and not VendorStatement.objects.exists()


# --------------------------------------------------------------------------- QA-032 / QA-033 / QA-009 / QA-034
# A lock is permanent, so the period must be named exactly and must be over.


@pytest.mark.parametrize("posted", ["", "   ", "not-a-date", "2026-13-45", "2026-9-30x"])
def test_lock_needs_an_exact_period_and_never_falls_back_to_the_last_month_end(client, org, approver, posted):
    login(client, approver)
    r = client.post(reverse("close:lock"), {"period": posted, "expected_version": "0", "note": ""}, follow=True)
    assert "Choose which period to lock" in r.content.decode()
    assert not AccrualSnapshot.objects.filter(organization=org).exists()


def test_lock_without_a_period_field_at_all_locks_nothing(client, org, approver):
    login(client, approver)
    client.post(reverse("close:lock"), {"expected_version": "0"})
    assert not AccrualSnapshot.objects.filter(organization=org).exists()


def test_a_period_cannot_be_locked_on_its_own_last_day(client, org, approver):
    login(client, approver)
    today = timezone.localdate()
    r = client.post(reverse("close:lock"), {"period": today.isoformat(), "expected_version": "0"}, follow=True)
    assert "hasn't ended yet" in r.content.decode()
    assert not AccrualSnapshot.objects.filter(organization=org).exists()
    yesterday = today - timedelta(days=1)
    client.post(reverse("close:lock"), {"period": yesterday.isoformat(), "expected_version": "0"})
    assert AccrualSnapshot.objects.filter(organization=org, period_end=yesterday).exists()


def test_the_lock_button_is_only_offered_for_a_period_that_has_ended(client, org, approver):
    login(client, approver)
    future = (timezone.localdate() + timedelta(days=40)).isoformat()
    page = client.get(reverse("close:accruals"), {"period": future, "live": "1"}).content.decode()
    assert "Lock the period ending" not in page and "hasn't ended yet, so it can't be locked" in page
    page = client.get(reverse("close:accruals"), {"period": "2026-09-30"}).content.decode()
    assert "Lock the period ending" in page


def test_a_second_version_is_only_offered_when_something_changed(client, org, approver):
    settings_for(org, expect_destination="never")
    s = shipment(org, ship="2026-09-12")
    invoice(org, s, day="2026-09-20")
    login(client, approver)
    client.post(reverse("close:lock"), {"period": "2026-09-30", "expected_version": "0", "note": ""})

    page = client.get(reverse("close:accruals"), {"period": "2026-09-30", "live": "1"}).content.decode()
    assert "Nothing has changed since it was locked" in page and "Lock the period ending" not in page

    invoice(org, s, day="2026-10-02", lines=(("Customs Clearance", "195.00"),))
    page = client.get(reverse("close:accruals"), {"period": "2026-09-30", "live": "1"}).content.decode()
    assert "Lock the period ending" in page


def test_a_garbled_period_in_the_address_says_so(client, org, approver):
    login(client, approver)
    r = client.get(reverse("close:accruals"), {"period": "abc"}, follow=True)
    assert "isn&#x27;t a date" in r.content.decode() or "isn't a date" in r.content.decode()


# --------------------------------------------------------------------------- QA-035 and amounts that can't be stored
# Typed money must fit the database column, and a currency must be exactly three letters (not cut to three).


@pytest.mark.parametrize("raw, expected", [
    ("1,250.5", Decimal("1250.50")), (" 640 ", Decimal("640.00")), ("0.005", Decimal("0.01")),
    ("99999999999999.99", Decimal("99999999999999.99")),
])
def test_parse_amount_accepts_what_fits(raw, expected):
    from apps.core.money import parse_amount

    assert parse_amount(raw) == expected


@pytest.mark.parametrize("raw, kind", [
    ("abc", "number"), ("", "number"), ("1e9", "number"), ("NaN", "number"), ("Infinity", "number"),
    ("-5", "number"), ("12abc", "number"),
    ("100000000000000", "range"), ("99999999999999999999", "range"),
])
def test_parse_amount_refuses_what_cannot_be_stored(raw, kind):
    from apps.core.money import AmountError, parse_amount

    with pytest.raises(AmountError) as e:
        parse_amount(raw)
    assert e.value.kind == kind


def test_a_currency_longer_than_three_letters_is_refused_not_cut(client, org, approver):
    login(client, approver)
    r = client.post(reverse("close:payment_add"), {"vendor": HL, "paid_on": "2026-09-01", "amount": "10",
                                                   "currency": "DOLLARS"}, follow=True)
    assert "three-letter code" in r.content.decode()
    assert not VendorPayment.objects.exists()


@pytest.mark.parametrize("amount", ["99999999999999999999", "1e30", "100000000000000"])
def test_an_overflowing_payment_is_refused_and_pages_still_load(client, org, approver, amount):
    login(client, approver)
    r = client.post(reverse("close:payment_add"), {"vendor": HL, "paid_on": "2026-09-01", "amount": amount,
                                                   "currency": "USD"}, follow=True)
    assert "too large" in r.content.decode() or "as a number" in r.content.decode()
    assert not VendorPayment.objects.exists()
    assert client.get(reverse("close:statements")).status_code == 200


@pytest.mark.parametrize("amount", ["99999999999999999999", "100000000000000"])
def test_an_overflowing_adjustment_is_refused(client, org, approver, amount):
    s = shipment(org, ship="2026-09-12")
    login(client, approver)
    r = client.post(reverse("close:adjust"), {"shipment": s.pk, "group": "destination", "action": "amount",
                                              "amount": amount, "note": "typo", "period": "2026-09-30"}, follow=True)
    assert "too large" in r.content.decode()
    assert not AccrualAdjustment.objects.exists()
    assert client.get(reverse("close:accruals"), {"period": "2026-09-30", "live": "1"}).status_code == 200


def test_an_overflowing_statement_balance_is_refused(client, org, approver):
    st = _statement(org, [("invoice", "HA-1", date(2026, 8, 20), "", "100.00")])
    login(client, approver)
    r = client.post(reverse("close:statement_edit", args=[st.pk]), {"vendor": HL, "statement_date": "2026-09-30",
                                                                     "currency": "USD",
                                                                     "closing_balance": "99999999999999999999"},
                    follow=True)
    assert "too large" in r.content.decode()
    st.refresh_from_db()
    assert st.closing_balance is None or abs(st.closing_balance) < Decimal("1e14")
    assert client.get(reverse("close:statement", args=[st.pk])).status_code == 200


def test_an_overflowing_approval_limit_is_refused(client, admin_user, org, user):
    from apps.core.models import Membership

    client.force_login(admin_user)
    m = Membership.objects.get(user=user, organization=org)
    r = client.post(reverse("core:update_member", args=[m.pk]), {"role": "approver",
                                                                  "approval_limit": "99999999999999999999"},
                    follow=True)
    assert "too large" in r.content.decode()
    m.refresh_from_db()
    assert m.role == "reviewer" and m.approval_limit is None
