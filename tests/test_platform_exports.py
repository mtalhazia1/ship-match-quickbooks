"""Exports: shipments with totals and issues, documents with every value read, validation issues; CSV streamed
and Excel; list-page filters; the formula guard; permissions; audit; the API with scopes."""
import csv
import io

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.services.apikeys import create_key
from apps.core.models import AuditEvent, Membership, Organization
from apps.documents.models import ExtractedField
from apps.documents.services.ingest import ingest_bytes
from apps.shipments.models import Shipment

URL = "integrations:export"


@pytest.fixture
def loaded(org, dataset):
    for f in ["S03_1_commercial_invoice.pdf", "S03_2_bill_of_lading.pdf", "S03_3_freight_invoice.pdf",
              "S01_1_commercial_invoice.pdf", "S01_2_bill_of_lading.pdf", "S01_3_freight_invoice.pdf"]:
        ingest_bytes(org, f, (dataset / "pdf" / f).read_bytes(), process="sync")
    return org


def _csv(response) -> list[dict]:
    body = b"".join(response.streaming_content).decode("utf-8")
    assert body.startswith("﻿")
    return list(csv.DictReader(io.StringIO(body.lstrip("﻿"))))


def _xlsx(response) -> list[list]:
    wb = load_workbook(io.BytesIO(b"".join(response.streaming_content)))
    return [[c.value for c in row] for row in wb.active.iter_rows()]


@pytest.mark.django_db
def test_shipments_csv_has_totals_issues_and_follows_the_list_filters(client, viewer, loaded):
    client.force_login(viewer)
    r = client.get(reverse(URL, args=["shipments"]), {"status": "all", "format": "csv"})
    assert r.status_code == 200 and r["Content-Type"].startswith("text/csv")
    assert r.streaming and 'attachment; filename="shipmatch-shipments-test-' in r["Content-Disposition"]
    rows = _csv(r)
    assert len(rows) == Shipment.objects.filter(organization=loaded).count() == 2
    by_ref = {row["Shipment"]: row for row in rows}
    bad = Shipment.objects.get(organization=loaded, issues__code="total_mismatch")
    row = by_ref[bad.reference]
    assert row["Bill of lading"] == bad.bl_number and row["Status"] == "Needs review"
    assert row["Documents"] == "3" and row["Invoices"] == "2"
    assert row["Total by currency"].startswith(("USD ", "EUR ")) and row["Open errors"] != "0"
    assert "Total doesn't match line items" in row["Open issues"]
    assert row["Link"].endswith(f"/review/shipments/{bad.pk}/")
    # Same filters as the queue: the default tab is "needs review"; search narrows it.
    rows = _csv(client.get(reverse(URL, args=["shipments"]), {"format": "csv"}))
    assert [r_["Shipment"] for r_ in rows] == [bad.reference]
    good = Shipment.objects.exclude(pk=bad.pk).get(organization=loaded)
    rows = _csv(client.get(reverse(URL, args=["shipments"]), {"format": "csv", "status": "all", "q": good.bl_number}))
    assert [r_["Shipment"] for r_ in rows] == [good.reference]
    event = AuditEvent.objects.filter(organization=loaded, action="export.downloaded").first()
    assert event.actor == viewer and event.data["kind"] == "shipments" and event.data["filters"]["q"] == good.bl_number


@pytest.mark.django_db
def test_documents_export_has_every_value_read(client, user, loaded):
    client.force_login(user)
    rows = _csv(client.get(reverse(URL, args=["documents"]), {"view": "all", "format": "csv"}))
    assert len(rows) == 6
    invoice = next(r for r in rows if r["Type"] == "Freight invoice")
    for column in ("Invoice number", "Total amount", "Currency", "B/L number", "Container numbers", "Vendor name"):
        assert column in invoice
    assert invoice["Invoice number"] and invoice["Shipment"].startswith("SHP-")
    bl = next(r for r in rows if r["Type"] == "Bill of lading")
    assert bl["Port of loading"] and bl["Invoice number"] == ""
    # The default view is "not in a shipment", as on the Documents page.
    assert _csv(client.get(reverse(URL, args=["documents"]), {"format": "csv"})) == []
    rows = _csv(client.get(reverse(URL, args=["documents"]), {"view": "all", "type": "bill_of_lading",
                                                              "format": "csv"}))
    assert {r["Type"] for r in rows} == {"Bill of lading"} and len(rows) == 2


@pytest.mark.django_db
def test_issues_export(client, viewer, loaded):
    client.force_login(viewer)
    rows = _csv(client.get(reverse(URL, args=["issues"]), {"status": "all", "format": "csv"}))
    assert rows and {"Shipment", "Severity", "Issue", "Money at risk", "Open"} <= set(rows[0])
    assert any(r["Issue"] == "Total doesn't match line items" and r["Severity"] == "Error" for r in rows)
    errors = _csv(client.get(reverse(URL, args=["issues"]), {"status": "all", "severity": "error", "format": "csv"}))
    assert errors and all(r["Severity"] == "Error" for r in errors)


@pytest.mark.django_db
def test_formula_guard_in_csv_and_excel(client, viewer, loaded):
    doc = loaded.documents.filter(doc_type="freight_invoice").first()
    ExtractedField.objects.filter(document=doc, name="vendor_name").update(value='=HYPERLINK("http://evil","x")')
    doc.original_filename = "+cmd|' /C calc'!A0.pdf"
    doc.save(update_fields=["original_filename"])
    client.force_login(viewer)
    rows = _csv(client.get(reverse(URL, args=["documents"]), {"view": "all", "format": "csv"}))
    row = next(r for r in rows if r["Document id"] == str(doc.pk))
    assert row["Vendor name"] == "'=HYPERLINK(\"http://evil\",\"x\")" and row["File"].startswith("'+cmd")
    assert row["Total amount"] and not row["Total amount"].startswith("'")    # numbers stay numbers

    r = client.get(reverse(URL, args=["documents"]), {"view": "all", "format": "xlsx"})
    assert r.status_code == 200 and r["Content-Type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(b"".join(r.streaming_content)))
    ws = wb.active
    header = [c.value for c in ws[1]]
    for row_cells in ws.iter_rows(min_row=2):
        if row_cells[0].value == doc.pk:
            vendor = row_cells[header.index("Vendor name")]
            assert vendor.data_type == "s" and vendor.value.startswith("'=")
            assert isinstance(row_cells[header.index("Total amount")].value, float)
            break
    else:
        raise AssertionError("document row missing")


@pytest.mark.django_db
def test_shipments_excel(client, viewer, loaded):
    client.force_login(viewer)
    r = client.get(reverse(URL, args=["shipments"]), {"status": "all", "format": "xlsx"})
    rows = _xlsx(r)
    assert rows[0][0] == "Shipment" and len(rows) == 3
    assert {row[0] for row in rows[1:]} == set(Shipment.objects.filter(organization=loaded)
                                                .values_list("reference", flat=True))


@pytest.mark.django_db
def test_exports_need_sign_in_and_stay_in_the_organization(client, loaded, viewer, dataset):
    r = client.get(reverse(URL, args=["shipments"]), {"status": "all", "format": "csv"})
    assert r.status_code == 302 and reverse("accounts:login") in r.url
    other = Organization.objects.create(name="Other Co", slug="other")
    ingest_bytes(other, "x.pdf", (dataset / "pdf" / "S05_2_bill_of_lading.pdf").read_bytes(), process="sync")
    outsider = get_user_model().objects.create_user("outsider", "o@x.example", "pw-123456789-test")
    Membership.objects.create(user=outsider, organization=other, role="viewer")
    client.force_login(outsider)
    rows = _csv(client.get(reverse(URL, args=["shipments"]), {"status": "all", "format": "csv", "org": "test"}))
    refs = {r["Shipment"] for r in rows}
    assert refs == set(Shipment.objects.filter(organization=other).values_list("reference", flat=True))


@pytest.mark.django_db
def test_bad_export_requests_explain_themselves(client, viewer, loaded):
    client.force_login(viewer)
    r = client.get(reverse(URL, args=["shipments"]), {"format": "pdf"}, follow=True)
    assert b"Choose CSV or Excel" in r.content
    r = client.get(reverse(URL, args=["shipments"]), {"format": "csv", "status": "lost"}, follow=True)
    assert b"Choose a shipment status" in r.content
    r = client.get(reverse(URL, args=["documents"]), {"format": "csv", "from": "31/12/2026"}, follow=True)
    assert b"isn&#x27;t a date" in r.content or b"isn't a date" in r.content
    assert client.get(reverse(URL, args=["secrets"])).status_code == 404


@pytest.mark.django_db
def test_export_menu_on_the_list_pages_keeps_the_filters(client, viewer, loaded):
    client.force_login(viewer)
    html = client.get(reverse("review:queue"), {"status": "all", "q": "MSCU"}).content.decode()
    assert reverse(URL, args=["shipments"]) + "?status=all&amp;q=MSCU&amp;format=csv" in html
    assert reverse(URL, args=["issues"]) in html
    html = client.get(reverse("review:documents"), {"view": "matched"}).content.decode()
    assert reverse(URL, args=["documents"]) + "?view=matched&amp;format=xlsx" in html


@pytest.mark.django_db
def test_api_exports_need_the_exports_scope(client, loaded):
    _key, token = create_key(loaded, "Warehouse", "viewer", None, None, scopes=["exports:read"])
    r = client.get(f"/api/{loaded.slug}/exports/shipments", {"status": "all"}, HTTP_AUTHORIZATION=f"Bearer {token}")
    assert r.status_code == 200 and len(_csv(r)) == 2
    r = client.get(f"/api/{loaded.slug}/exports/issues", {"status": "all", "format": "xlsx", "open": "true"},
                   HTTP_AUTHORIZATION=f"Bearer {token}")
    assert r.status_code == 200 and _xlsx(r)[0][0] == "Shipment"
    event = AuditEvent.objects.filter(action="export.downloaded").first()
    assert event.data["api_key"] == "Warehouse" and event.actor is None
    _key, narrow = create_key(loaded, "Reader", "viewer", None, None, scopes=["shipments:read"])
    r = client.get(f"/api/{loaded.slug}/exports/shipments", HTTP_AUTHORIZATION=f"Bearer {narrow}")
    assert r.status_code == 403 and "exports:read" in r.json()["detail"]
    r = client.get(f"/api/{loaded.slug}/exports/documents", {"view": "nope"}, HTTP_AUTHORIZATION=f"Bearer {token}")
    assert r.status_code == 400
