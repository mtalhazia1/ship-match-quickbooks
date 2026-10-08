"""Landed cost per product and freight invoices shared by several shipments (apps/landed).

Covers: the cent-exact split math, every allocation basis (value, quantity, weight, volume, per charge type,
shipment override, source hints), currency conversion, fallback notes, charge sources, freezing at approval,
CSV/XLSX exports, the per-product report, shared-invoice detection with and without lines that name a
shipment, manual splits, confirmation, approval and posting rules, permissions and the synthetic helpers.
"""
from __future__ import annotations

import csv
import io
import itertools
import uuid
from datetime import date
from decimal import Decimal as D

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse

from apps.core.models import AuditEvent, Membership, Organization
from apps.documents.models import Document, ExtractedField
from apps.documents.services.corrections import after_correction, correct_field
from apps.documents.services.ingest import ingest_bytes
from apps.landed.models import (
    InvoiceAllocation,
    LandedCostRun,
    LandedSettings,
    SharedInvoice,
    ShipmentLandedOverride,
)
from apps.landed.services import allocation, exports, landed, rounding
from apps.landed.services import charges as lc_charges
from apps.landed.services import report as lc_report
from apps.shipments.models import MatchLink, Shipment, ValidationIssue
from apps.shipments.services.approval import approval_blockers, posting_blockers
from apps.shipments.services.containers import make_container
from apps.shipments.services.validation import validate_shipment
from synthetic import landed as synth
from tests.conftest import PASSWORD

_seq = itertools.count(1)
APEX = "Apex Housewares Manufacturing Ltd."
HL = "Harborlink Logistics LLC"


# --------------------------------------------------------------------------- helpers


def _doc(org, shipment, doc_type, fields, text="", method=MatchLink.Method.EXACT_BL):
    n = next(_seq)
    d = Document.objects.create(organization=org, original_filename=f"{doc_type}-{n}.pdf", sha256=uuid.uuid4().hex,
                                doc_type=doc_type, text=text,
                                status=Document.Status.MATCHED if shipment else Document.Status.UNMATCHED)
    for k, v in fields.items():
        ExtractedField.objects.create(document=d, name=k, value=v, confidence=0.99, grounded=True)
    if shipment is not None:
        MatchLink.objects.create(document=d, shipment=shipment, method=method, score=1.0)
    return d


def product(desc, qty=None, amount="0.00", **extra):
    item = {"description": desc, "amount": amount}
    if qty is not None:
        item["quantity"] = str(qty)
    item.update({k: str(v) for k, v in extra.items()})
    return item


def make_shipment(org, products, charges=(("Ocean Freight", "1000.00"),), *, ci_currency="USD", fr_currency="USD",
                  boxes=1, invoice_date="2026-03-05", vendor=APEX):
    n = next(_seq)
    containers = [make_container("OSL", 200000 + n * 13 + i) for i in range(boxes)]
    bl = f"OSLN{8000000000 + n * 7919}"
    s = Shipment.objects.create(organization=org, bl_number=bl, container_numbers=containers, po_numbers=[f"PO-77{n:03d}"])
    _doc(org, s, "bill_of_lading", {"bl_number": bl, "container_numbers": containers, "issue_date": "2026-03-10"},
         text=f"BILL OF LADING B/L No.: {bl} " + " ".join(containers))
    total = sum((D(p["amount"]) for p in products), D("0.00"))
    ci = _doc(org, s, "commercial_invoice", {"vendor_name": vendor, "invoice_number": f"CI-{n}",
                                             "invoice_date": invoice_date, "currency": ci_currency,
                                             "line_items": products, "total_amount": f"{total:.2f}"})
    fr = None
    if charges:
        lines = [{"description": d, "amount": a} for d, a in charges]
        fr_total = sum((D(a) for _, a in charges), D("0.00"))
        fr = _doc(org, s, "freight_invoice", {"vendor_name": HL, "invoice_number": f"FR-{n}", "currency": fr_currency,
                                              "bl_number": bl, "container_numbers": containers, "line_items": lines,
                                              "total_amount": f"{fr_total:.2f}"},
                  text=f"FREIGHT INVOICE B/L No.: {bl} Container(s): {', '.join(containers)}")
    return s, ci, fr


def parts(lc, category="freight"):
    return [p.part(category) for p in lc.products]


def ingest_all(org, files):
    for name, data in files:
        ingest_bytes(org, name, data, process="sync")


@pytest.fixture
def shared(org):
    """Three shipments and a forwarder invoice that names each B/L on its lines (the demo scenario)."""
    sc = synth.scenario(seed=7)
    ingest_all(org, sc["files"])
    doc = Document.objects.get(organization=org, doc_type="freight_invoice")
    ships = [Shipment.objects.get(organization=org, bl_number=s.bl) for s in sc["shipments"]]
    return {"doc": doc, "ships": ships, "sc": sc}


def approve(client, approver, shipment):
    client.force_login(approver)
    return client.post(reverse("review:approve", args=[shipment.pk]))


# --------------------------------------------------------------------------- split math


def test_split_adds_up_exactly_and_hands_out_leftover_cents():
    assert rounding.split(10000, [1, 1, 1]) == [3334, 3333, 3333]
    assert sum(rounding.split(10001, [D("0.3"), D("0.3"), D("0.4")])) == 10001
    assert rounding.split(100, [2, 1, 0]) == [67, 33, 0]
    assert rounding.split(-10000, [1, 1, 1]) == [-3334, -3333, -3333]   # credits split the same way
    assert rounding.split(5, [1, 1, 1, 1, 1, 1, 1]) == [1, 1, 1, 1, 1, 0, 0]
    with pytest.raises(ValueError):
        rounding.split(100, [0, 0])
    with pytest.raises(ValueError):
        rounding.split(100, [1, -1])


@pytest.mark.parametrize("columns,targets", [
    ([10000, 3333, 1], [5000, 8334]), ([999, 1, 1, 1], [1, 1000, 1]), ([-500, 2500], [1000, 1000]),
])
def test_split_table_keeps_rows_and_columns_exact(columns, targets):
    t = rounding.split_table(columns, targets)
    assert [sum(r) for r in t] == targets
    assert [sum(t[r][c] for r in range(len(targets))) for c in range(len(columns))] == columns


# --------------------------------------------------------------------------- landed cost: bases


def two_products(**weights):
    return [product("Steel bracket", 10, "1000.00", sku="BRK-1", weight_kg=weights.get("w1", "300"),
                    volume_cbm=weights.get("v1", "1")),
            product("Steel hinge", 30, "3000.00", sku="HNG-2", weight_kg=weights.get("w2", "100"),
                    volume_cbm=weights.get("v2", "3"))]


@pytest.mark.parametrize("method,expected", [
    ("value", ["100.00", "300.00"]), ("quantity", ["100.00", "300.00"]), ("weight", ["300.00", "100.00"]),
    ("volume", ["100.00", "300.00"]),
])
def test_each_method_spreads_the_charge(org, method, expected):
    s, _, _ = make_shipment(org, two_products(), charges=[("Ocean Freight", "400.00")])
    LandedSettings.objects.create(organization=org, method=method)
    lc = landed.compute(s)
    assert lc.complete
    assert [str(x) for x in parts(lc)] == expected
    assert lc.charges_total == D("400.00") and lc.landed_total == D("4400.00")
    assert [r.used for r in lc.charges] == [method]


def test_rounding_keeps_every_charge_exact_and_per_unit_and_uplift(org):
    items = [product(f"Mug {i}", 3, "100.00", sku=f"M{i}") for i in range(3)]
    s, _, _ = make_shipment(org, items, charges=[("Ocean Freight", "100.00"), ("Terminal Handling Charge", "0.01")])
    lc = landed.compute(s)
    assert parts(lc) == [D("33.34"), D("33.33"), D("33.33")]
    assert parts(lc, "handling") == [D("0.01"), D("0.00"), D("0.00")]
    assert sum(p.charges for p in lc.products) == D("100.01")
    first = lc.products[0]
    assert first.landed == D("133.35") and first.per_unit == D("44.4500") and first.uplift == D("33.4")
    assert lc.uplift == D("33.3")


def test_per_charge_type_with_source_hint_and_shipment_override(org, monkeypatch):
    def customs(shipment):
        yield {"code": "customs_duty", "amount": "200.00", "currency": "USD", "basis": "value",
               "description": "Duty", "source": "Customs entry 1"}

    monkeypatch.setattr(lc_charges, "CHARGE_SOURCES", [])
    lc_charges.register_charge_source(customs)
    lc_charges.register_charge_source(customs)   # registering twice adds it once
    assert lc_charges.CHARGE_SOURCES == [customs]
    s, _, _ = make_shipment(org, two_products(), charges=[("Ocean Freight", "400.00")])
    LandedSettings.objects.create(organization=org, method="quantity", by_category={"freight": "weight"})
    lc = landed.compute(s)
    assert parts(lc, "freight") == [D("300.00"), D("100.00")]       # freight by weight (organization, per type)
    assert parts(lc, "duty") == [D("50.00"), D("150.00")]           # duty by value (the source's hint)
    assert lc.categories == [("freight", "Freight"), ("duty", "Duties and taxes")]
    # A shipment's own setting wins over the organization's; the hint still applies to duty.
    ShipmentLandedOverride.objects.create(shipment=s, method="weight", by_category={})
    lc = landed.compute(s)
    assert lc.policy.source == "shipment"
    assert parts(lc, "freight") == [D("300.00"), D("100.00")] and parts(lc, "duty") == [D("50.00"), D("150.00")]
    ShipmentLandedOverride.objects.filter(shipment=s).update(by_category={"duty": "volume"})
    assert parts(landed.compute(s), "duty") == [D("50.00"), D("150.00")]   # volumes 1 and 3


def test_currency_conversion_and_missing_rate(org):
    org.fx_rates = {"EUR": "1.0837"}
    org.save()
    items = [product(f"Towel {i}", 10, "100.00") for i in range(3)]
    s, _, _ = make_shipment(org, items, ci_currency="EUR", charges=[("Ocean Freight", "100.00")], fr_currency="EUR")
    lc = landed.compute(s)
    assert [p.value_home for p in lc.products] == [D("108.37")] * 3
    assert lc.charges[0].amount_home == D("108.37")
    assert parts(lc) == [D("36.13"), D("36.12"), D("36.12")]
    assert lc.products[0].value == D("100.00") and lc.products[0].currency == "EUR"
    s2, _, _ = make_shipment(org, items, ci_currency="GBP")
    lc2 = landed.compute(s2)
    assert not lc2.complete and "No exchange rate for GBP" in lc2.notes[0].text
    assert all(not p.allocated for p in lc2.products)


def test_weight_basis_falls_back_to_value_with_a_visible_note(org):
    items = two_products()
    del items[1]["weight_kg"]
    s, _, _ = make_shipment(org, items, charges=[("Ocean Freight", "400.00"), ("THC", "40.00")])
    LandedSettings.objects.create(organization=org, method="weight")
    lc = landed.compute(s)
    assert parts(lc) == [D("100.00"), D("300.00")]
    assert lc.charges[0].used == "value" and lc.charges[0].fell_back
    texts = [n.text for n in lc.warnings]
    assert "Weight is missing on 1 of 2 products, so the freight and handling charges were spread by value instead " \
           "of by weight." in texts


def test_quantity_basis_without_quantities_and_zero_values_fall_back(org):
    items = [product("Sample A", None, "0.00", sku="S-A"), product("Sample B", None, "0.00", sku="S-B")]
    s, _, _ = make_shipment(org, items, charges=[("Ocean Freight", "10.01")])
    LandedSettings.objects.create(organization=org, method="quantity")
    lc = landed.compute(s)
    assert parts(lc) == [D("5.01"), D("5.00")] and lc.charges[0].used == "equal"
    assert any("spread equally instead of by quantity" in n.text for n in lc.warnings)
    assert lc.products[0].per_unit is None and lc.products[0].uplift is None


def test_missing_commercial_invoice_or_lines_are_explained(org):
    s = Shipment.objects.create(organization=org, bl_number="OSLN1234567890")
    _doc(org, s, "freight_invoice", {"vendor_name": HL, "invoice_number": "X1", "total_amount": "100.00",
                                     "line_items": [{"description": "Ocean Freight", "amount": "100.00"}]})
    lc = landed.compute(s)
    assert not lc.complete and lc.notes[0].text.startswith("No commercial invoice in this shipment yet")
    _doc(org, s, "commercial_invoice", {"vendor_name": APEX, "invoice_number": "CI-EMPTY", "total_amount": "50.00"})
    lc = landed.compute(s)
    assert any("Commercial invoice CI-EMPTY has no product lines" in n.text for n in lc.notes)


def test_charge_lines_on_the_commercial_invoice_are_spread_not_counted_as_products(org):
    items = two_products() + [product("Sea freight to Long Beach", None, "80.00"), product("Discount", None, "-40.00")]
    s, _, _ = make_shipment(org, items, charges=())
    lc = landed.compute(s)
    assert [p.description for p in lc.products] == ["Steel bracket", "Steel hinge"]
    assert lc.charges_total == D("40.00")
    assert any("read like a charge or discount" in n.text for n in lc.notes)


def test_credit_notes_reduce_and_possible_duplicates_are_left_out(org):
    s, _, fr = make_shipment(org, two_products(), charges=[("Ocean Freight", "400.00")])
    _doc(org, s, "credit_note", {"vendor_name": HL, "credit_note_number": "CN-1", "total_amount": "100.00",
                                 "line_items": [{"description": "Ocean freight overcharge", "amount": "100.00"}]})
    lc = landed.compute(s)
    assert lc.charges_total == D("300.00")
    copy = _doc(org, s, "freight_invoice", {"vendor_name": HL, "invoice_number": fr.field("invoice_number"),
                                            "total_amount": "400.00"})
    ValidationIssue.objects.create(organization=org, shipment=s, document=copy, code="duplicate_invoice",
                                   severity="error", message="dup", fingerprint="dup")
    lc = landed.compute(s)
    assert lc.charges_total == D("300.00")
    assert any("may be a duplicate" in n.text for n in lc.warnings)


def test_a_failing_charge_source_is_noted_and_the_others_still_count(org, monkeypatch):
    def broken(shipment):
        raise RuntimeError("customs service down")

    broken.label = "customs entries"
    monkeypatch.setattr(lc_charges, "CHARGE_SOURCES", [broken, lambda s: [{"code": "insurance", "amount": "12.00"}]])
    s, _, _ = make_shipment(org, two_products(), charges=[("Ocean Freight", "400.00")])
    lc = landed.compute(s)
    assert lc.charges_total == D("412.00")
    assert parts(lc, "insurance") == [D("3.00"), D("9.00")]
    assert any("Charges from customs entries couldn't be read" in n.text for n in lc.warnings)


def test_landed_cost_is_frozen_at_approval_and_released_on_reopen(client, org, approver):
    org.fx_rates = {"EUR": "1.10"}
    org.save()
    s, _, _ = make_shipment(org, two_products(), ci_currency="EUR", charges=[("Ocean Freight", "400.00")])
    assert approve(client, approver, s).status_code == 302
    s.refresh_from_db()
    assert s.status == "approved"
    run = LandedCostRun.objects.get(shipment=s)
    assert run.goods_total == D("4400.00") and run.lines.count() == 2
    org.fx_rates = {"EUR": "2.00"}
    org.save()
    s = Shipment.objects.select_related("organization").get(pk=s.pk)
    frozen = landed.landed_for(s)
    assert frozen.frozen_at and frozen.goods_total == D("4400.00") and frozen.charges_total == D("400.00")
    page = client.get(reverse("review:shipment", args=[s.pk])).content.decode()
    assert "Frozen when the shipment was approved" in page
    client.post(reverse("review:reopen", args=[s.pk]))
    assert not LandedCostRun.objects.filter(shipment=s).exists()
    assert landed.landed_for(Shipment.objects.get(pk=s.pk)).goods_total == D("8000.00")


# --------------------------------------------------------------------------- exports


def test_shipment_csv_and_xlsx_exports(client, org, viewer):
    items = two_products()
    items[0]["description"] = '=HYPERLINK("http://evil.example")'
    items[1]["sku"] = "00123"
    s, _, _ = make_shipment(org, items, charges=[("Ocean Freight", "400.00")])
    client.force_login(viewer)
    r = client.get(reverse("landed:shipment_export", args=[s.pk, "csv"]))
    assert r.status_code == 200 and r["Content-Disposition"].endswith(f'landed-cost-{s.reference}.csv"')
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][0] == "SKU" and "Landed cost per unit (USD)" in rows[0]
    assert rows[1][1] == "'=HYPERLINK(\"http://evil.example\")"
    assert rows[-1][1] == "Total" and rows[-1][rows[0].index("Landed cost (USD)")] == "4400.00"

    r = client.get(reverse("landed:shipment_export", args=[s.pk, "xlsx"]))
    assert r.status_code == 200 and r["Content-Type"] == exports.XLSX_TYPE
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Landed cost", "Charges", "Notes"]
    ws = wb["Landed cost"]
    head = [c.value for c in ws[3]]
    bad = ws.cell(row=4, column=2)
    assert bad.value == '=HYPERLINK("http://evil.example")' and bad.data_type == "s"
    assert ws.cell(row=5, column=1).value == "00123"
    landed_col = head.index("Landed cost (USD)") + 1
    assert ws.cell(row=4, column=landed_col).value == pytest.approx(1100.0)
    assert client.get(reverse("landed:shipment_export", args=[s.pk, "pdf"])).status_code == 404
    assert AuditEvent.objects.filter(action="landed.exported", actor=viewer).count() == 2


# --------------------------------------------------------------------------- report


def test_report_by_product_average_last_and_change(client, org, viewer, approver):
    a, _, _ = make_shipment(org, [product("Cookware set", 100, "4000.00", sku="APX-CW10")],
                            charges=[("Ocean Freight", "1000.00")], invoice_date="2026-03-01")
    b, _, _ = make_shipment(org, [product("Cookware set", 100, "4200.00", sku="APX-CW10")],
                            charges=[("Ocean Freight", "1300.00")], invoice_date="2026-04-01")
    open_one, _, _ = make_shipment(org, [product("Cookware set", 50, "2000.00", sku="APX-CW10")],
                                   charges=[("Ocean Freight", "100.00")], invoice_date="2026-04-10")
    for s in (a, b):
        Shipment.objects.filter(pk=s.pk).update(status="approved")
    year = (date(2026, 1, 1), date(2026, 12, 31))
    rep = lc_report.build(org, *year)
    assert LandedCostRun.objects.filter(shipment__in=[a, b]).count() == 2   # frozen on first view
    row = rep.rows[0]
    assert row.sku == "APX-CW10" and len(row.points) == 2
    assert row.last.per_unit == D("55.0000") and row.previous.per_unit == D("50.0000")
    assert row.average == D("52.5000") and row.change == D("10.0") and row.uplift == D("28.0")
    assert row.spark
    rep2 = lc_report.build(org, *year, include_open=True)
    assert len(rep2.rows[0].points) == 3 and rep2.in_review == 1 and rep2.rows[0].last.in_review
    assert lc_report.build(org, *year, q="nothing like it").rows == []

    client.force_login(viewer)
    page = client.get(reverse("landed:report") + "?period=all")
    assert page.status_code == 200 and "Cookware set" in page.content.decode() and "Up 10.0%" in page.content.decode()
    r = client.get(reverse("landed:report_export", args=["csv"]) + "?period=all&open=1")
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][:5] == ["SKU", "Description", "Supplier", "Date", "Shipment"] and len(rows) == 4
    assert [x[5] for x in rows[1:]] == ["Approved", "Approved", "In review"]
    r = client.get(reverse("landed:report_export", args=["xlsx"]) + "?period=all")
    assert r.status_code == 200 and r["Content-Type"] == exports.XLSX_TYPE
    client.logout()
    assert client.get(reverse("landed:report")).status_code == 302


# --------------------------------------------------------------------------- settings and permissions


def test_settings_need_manage_and_shipment_method_needs_edit(client, org, user, viewer, admin_user):
    s, _, _ = make_shipment(org, two_products(), charges=[("Ocean Freight", "400.00")])
    client.force_login(user)
    assert client.get(reverse("landed:settings")).status_code == 403
    r = client.post(reverse("landed:shipment_method", args=[s.pk]), {"method": "weight", "cat_duty": "value"})
    assert r.status_code == 302
    o = ShipmentLandedOverride.objects.get(shipment=s)
    assert o.method == "weight" and o.by_category == {"duty": "value"}
    assert AuditEvent.objects.filter(action="landed.method_changed", actor=user).exists()
    client.post(reverse("landed:shipment_method", args=[s.pk]), {"action": "clear"})
    assert not ShipmentLandedOverride.objects.filter(shipment=s).exists()
    assert client.post(reverse("landed:shipment_method", args=[s.pk]), {"method": "colour"}).status_code == 302
    assert not ShipmentLandedOverride.objects.filter(shipment=s).exists()

    client.force_login(viewer)
    assert client.post(reverse("landed:shipment_method", args=[s.pk]), {"method": "weight"}).status_code == 403
    assert client.get(reverse("review:shipment", args=[s.pk])).status_code == 200

    client.force_login(admin_user)
    assert client.get(reverse("landed:settings")).status_code == 200
    r = client.post(reverse("landed:settings"), {"method": "volume", "cat_freight": "weight", "cat_duty": "volume"})
    assert r.status_code == 302
    st = LandedSettings.objects.get(organization=org)
    assert st.method == "volume" and st.by_category == {"freight": "weight"}   # same as default is not stored
    assert AuditEvent.objects.filter(action="landed.settings_updated").exists()

    Shipment.objects.filter(pk=s.pk).update(status="approved")
    client.force_login(user)
    client.post(reverse("landed:shipment_method", args=[s.pk]), {"method": "weight"})
    assert not ShipmentLandedOverride.objects.filter(shipment=s).exists()   # locked: frozen at approval


def test_other_organizations_are_out_of_reach(client, org, user):
    other = Organization.objects.create(name="Other", slug="other")
    s, _, fr = make_shipment(other, two_products())
    client.force_login(user)
    assert client.get(reverse("landed:shipment_export", args=[s.pk, "csv"])).status_code == 404
    assert client.post(reverse("landed:split_confirm", args=[fr.pk])).status_code == 404
    assert client.post(reverse("landed:shipment_method", args=[s.pk]), {"method": "value"}).status_code == 404


# --------------------------------------------------------------------------- shared invoices: detection


def test_detects_a_shared_invoice_whose_lines_name_each_bl(shared):
    doc, ships = shared["doc"], shared["ships"]
    si = SharedInvoice.objects.get(document=doc)
    assert si.basis == "lines" and si.status == "active"
    rows = {r.shipment_id: r for r in allocation.allocations(doc)}
    assert set(rows) == {s.pk for s in ships}
    # Each leg's freight and handling lines go to it whole; Documentation and Customs (250.00) by containers 1:2:1.
    assert [rows[s.pk].amount for s in ships] == [D("2522.50"), D("5570.00"), D("2522.50")]
    assert sum(r.amount for r in rows.values()) == D(shared["sc"]["invoice"]["total_amount"])
    assert doc.match.shipment_id == ships[0].pk   # the invoice stays matched to one shipment
    assert "B/L " + ships[1].bl_number in rows[ships[1].pk].reason
    assert AuditEvent.objects.filter(action="shared_invoice.detected", object_id=str(doc.pk)).exists()
    # Every shipment is told about its share and asked to confirm the split.
    for s in ships:
        assert s.issues.filter(code="shared_invoice_split", resolved=False).exists()
    lines = allocation.bill_lines(doc)
    assert sum(x["amount"] for x in lines) == D(shared["sc"]["invoice"]["total_amount"])
    assert {x["shipment_reference"] for x in lines} == {s.reference for s in ships}
    demurrage = [x for x in lines if x["description"].startswith("Demurrage")]
    assert len(demurrage) == 1 and demurrage[0]["shipment_reference"] == ships[1].reference


def test_detects_a_shared_invoice_without_lines_that_name_a_shipment(org):
    sc = synth.scenario(seed=21, attributed=False)
    ingest_all(org, sc["files"])
    doc = Document.objects.get(organization=org, doc_type="freight_invoice")
    si = SharedInvoice.objects.get(document=doc)
    assert si.basis == "containers"
    total = D(sc["invoice"]["total_amount"])
    amounts = [r.amount for r in allocation.allocations(doc)]
    assert amounts == rounding.split_amount(total, [1, 2, 1]) and sum(amounts) == total


def test_a_bl_arriving_after_the_invoice_gets_its_share(org):
    sc = synth.scenario(seed=33, legs=2)
    files = dict(sc["files"])
    order = ["L33_1_bill_of_lading.pdf", "L33_1_commercial_invoice.pdf", "L33_shared_freight_invoice.pdf",
             "L33_2_bill_of_lading.pdf", "L33_2_commercial_invoice.pdf"]
    ingest_all(org, [(n, files[n]) for n in order[:3]])
    doc = Document.objects.get(organization=org, doc_type="freight_invoice")
    assert not InvoiceAllocation.objects.filter(document=doc).exists()
    ingest_all(org, [(n, files[n]) for n in order[3:]])
    rows = allocation.allocations(doc)
    assert len(rows) == 2 and sum(r.amount for r in rows) == D(sc["invoice"]["total_amount"])
    primary = doc.match.shipment
    assert primary.issues.filter(code="shared_invoice_split", resolved=False).exists()


def test_references_the_primary_also_has_dont_make_an_invoice_shared(org):
    a, _, fr = make_shipment(org, two_products())
    b = Shipment.objects.create(organization=org, bl_number="OSLN5550001111", po_numbers=list(a.po_numbers),
                                container_numbers=[a.container_numbers[0]])   # split PO, reused container
    ExtractedField.objects.create(document=fr, name="po_numbers", value=list(a.po_numbers))
    assert allocation.sync(fr) == set()
    assert not SharedInvoice.objects.filter(document=fr).exists()
    assert b.invoice_shares.count() == 0


def test_an_invoice_naming_an_approved_shipment_warns_and_shares_after_reopen(client, org, approver):
    a, _, fr = make_shipment(org, two_products(), charges=[("Ocean Freight", "1000.00")])
    b, _, _ = make_shipment(org, two_products(), charges=())
    Shipment.objects.filter(pk=b.pk).update(status="approved")
    fr.text += f" Also covers B/L: {b.bl_number}"
    fr.save()
    allocation.sync(fr)
    validate_shipment(a)
    issue = a.issues.get(code="shared_invoice_locked_ref")
    assert b.reference in issue.message and not InvoiceAllocation.objects.filter(document=fr).exists()
    client.force_login(approver)
    client.post(reverse("review:reopen", args=[b.pk]))
    assert {r.shipment_id for r in allocation.allocations(fr)} == {a.pk, b.pk}


# --------------------------------------------------------------------------- shared invoices: reviewers


def test_reviewer_types_a_manual_split_that_must_add_up(client, shared, user):
    doc, ships = shared["doc"], shared["ships"]
    total = D(shared["sc"]["invoice"]["total_amount"])
    url = reverse("landed:split_save", args=[doc.pk])
    client.force_login(user)
    short = {"basis": "manual", f"amount_{ships[0].pk}": "1000.00", f"amount_{ships[1].pk}": "1000.00",
             f"amount_{ships[2].pk}": "1000.00"}
    r = client.post(url, short, follow=True)
    assert "but the invoice total is" in r.content.decode() and "short by" in r.content.decode()
    assert SharedInvoice.objects.get(document=doc).basis == "lines"
    zero = {**short, f"amount_{ships[2].pk}": "0", f"amount_{ships[1].pk}": f"{total - 1000:.2f}"}
    assert "must be above zero" in client.post(url, zero, follow=True).content.decode()
    ok = {"basis": "manual", f"amount_{ships[0].pk}": "1,240.00", f"amount_{ships[1].pk}": f"{total - 2480:.2f}",
          f"amount_{ships[2].pk}": "1240.00"}
    r = client.post(url, ok)
    assert r.status_code == 302 and r["Location"].endswith("#shared-invoices")
    rows = {r.shipment_id: r.amount for r in allocation.allocations(doc)}
    assert rows == {ships[0].pk: D("1240.00"), ships[1].pk: total - 2480, ships[2].pk: D("1240.00")}
    si = SharedInvoice.objects.get(document=doc)
    assert si.basis == "manual" and allocation.is_confirmed(si) and si.confirmed_by == user
    ev = AuditEvent.objects.get(action="shared_invoice.split_changed")
    assert ev.actor == user and ev.data["after"][ships[0].reference] == "1240.00"
    # Bill lines follow the typed shares exactly.
    lines = allocation.bill_lines(doc)
    for s in ships:
        assert sum(x["amount"] for x in lines if x["shipment_reference"] == s.reference) == rows[s.pk]
    # The secondary shipment's page says what its share is, with links to the others.
    page = client.get(reverse("review:shipment", args=[ships[2].pk])).content.decode()
    assert "Shared invoice: your share" in page and "USD 1,240.00" in page and f"USD {total:,.2f}" in page
    assert reverse("review:shipment", args=[ships[0].pk]) in page
    # Re-detection never overrides the reviewer's split.
    allocation.sync(doc)
    assert {r.shipment_id: r.amount for r in allocation.allocations(doc)} == rows


@pytest.mark.parametrize("basis,expected_basis", [("equal", "equal"), ("containers", "containers"),
                                                  ("weight", "weight"), ("volume", "volume")])
def test_split_bases(shared, user, basis, expected_basis):
    doc = shared["doc"]
    total = D(shared["sc"]["invoice"]["total_amount"])
    result = allocation.save_split(doc, user, basis)
    assert result.basis == expected_basis and sum(result.shares) == rounding.to_cents(total)
    amounts = [r.amount for r in allocation.allocations(doc)]
    if basis == "equal":
        assert amounts == rounding.split_amount(total, [1, 1, 1])
    elif basis == "containers":
        assert amounts == rounding.split_amount(total, [1, 2, 1])
    else:
        key = "weight_kg" if basis == "weight" else "volume_cbm"
        measures = [sum((getattr(li, key) for li in s.lines), D(0)) for s in shared["sc"]["shipments"]]
        assert amounts == rounding.split_amount(total, measures)


def test_weight_split_without_weights_falls_back_to_containers(org, user):
    sc = synth.scenario(seed=44, legs=2)
    ingest_all(org, sc["files"])
    doc = Document.objects.get(organization=org, doc_type="freight_invoice")
    ci = Document.objects.filter(organization=org, doc_type="commercial_invoice").last()
    f = ci.fields.get(name="line_items")
    f.value = [{k: v for k, v in li.items() if k != "weight_kg"} for li in f.value]
    f.save()
    result = allocation.save_split(doc, user, "weight")
    assert result.basis == "containers" and "Split by containers instead of weight" in result.notes[0]


def test_confirming_resolves_the_split_warnings_and_the_other_shipments_containers(client, shared, user):
    doc, ships = shared["doc"], shared["ships"]
    primary = ships[0]
    open_container_errors = primary.issues.filter(code="container_not_on_bl", resolved=False)
    assert open_container_errors.count() == 3   # the invoice lists SHP-2's and SHP-3's containers
    assert any("Confirm the split of shared invoice" in b for b in approval_blockers(ships[1], user))
    client.force_login(user)
    r = client.post(reverse("landed:split_confirm", args=[doc.pk]))
    assert r.status_code == 302
    si = SharedInvoice.objects.get(document=doc)
    assert allocation.is_confirmed(si) and si.confirmed_by == user
    assert not primary.issues.filter(code="container_not_on_bl", resolved=False).exists()
    note = primary.issues.filter(code="container_not_on_bl").first().resolution_note
    assert ships[1].reference in note or ships[2].reference in note
    for s in ships:
        assert not s.issues.filter(code="shared_invoice_split", resolved=False).exists()
        assert not any("Confirm the split" in b for b in approval_blockers(s, user))
    assert AuditEvent.objects.filter(action="shared_invoice.share_confirmed", object_id=str(ships[1].pk)).exists()
    # A change to the invoice makes the confirmation stale again.
    total = doc.fields.get(name="total_amount")
    total.value = f"{D(total.value) + 100:.2f}"
    total.save()
    allocation.sync(doc)
    assert not allocation.is_confirmed(SharedInvoice.objects.get(document=doc))


def test_approval_order_and_posting_wait_for_every_shipment(client, shared, user, approver):
    doc, (primary, second, third) = shared["doc"], shared["ships"]
    allocation.confirm(doc, user)
    blockers = approval_blockers(primary, approver)
    assert any(f"Approve {second.reference} and {third.reference} first" in b for b in blockers)
    # Posting is blocked while any shipment with a share isn't approved.
    Shipment.objects.filter(pk=primary.pk).update(status="approved")
    primary.refresh_from_db()
    assert posting_blockers(primary) == [f"Invoice {allocation.invoice_number(doc)} is shared with {second.reference} "
                                         f"and {third.reference}, which must be approved before it is posted."]
    client.force_login(approver)
    r = client.post(reverse("review:post", args=[primary.pk]), follow=True)
    assert "Not posted." in r.content.decode()
    assert not AuditEvent.objects.filter(action="shipment.post_requested").exists()
    Shipment.objects.filter(pk=primary.pk).update(status="ready")
    for s in (second, third):
        assert approve(client, approver, s).status_code == 302
        s.refresh_from_db()
        assert s.status == "approved", approval_blockers(s, approver)
    primary.refresh_from_db()
    assert not any("first" in b for b in approval_blockers(primary, approver))
    # Once a shipment in the split is approved, the split is locked.
    with pytest.raises(allocation.SplitError, match="locked"):
        allocation.save_split(doc, user, "equal")
    Shipment.objects.filter(pk=primary.pk).update(status="approved")
    primary.refresh_from_db()
    assert posting_blockers(primary) == []
    # Reopening a shipment with a share stops posting again.
    client.post(reverse("review:reopen", args=[third.pk]))
    assert posting_blockers(primary)


def test_not_shared_and_back_to_automatic(client, shared, user):
    doc, ships = shared["doc"], shared["ships"]
    client.force_login(user)
    url = reverse("landed:split_dismiss", args=[doc.pk])
    assert "Say why" in client.post(url, {"note": ""}, follow=True).content.decode()
    client.post(url, {"note": "Other B/Ls are only quoted for reference."})
    si = SharedInvoice.objects.get(document=doc)
    assert si.status == "dismissed" and not InvoiceAllocation.objects.filter(document=doc).exists()
    assert allocation.bill_lines(doc) is None
    assert not ships[1].issues.filter(code__startswith="shared_", resolved=False).exists()
    allocation.sync(doc)   # detection respects the decision
    assert not InvoiceAllocation.objects.filter(document=doc).exists()
    client.post(reverse("landed:split_reset", args=[doc.pk]))
    assert len(allocation.allocations(doc)) == 3
    assert SharedInvoice.objects.get(document=doc).basis == "lines"


def test_start_a_split_the_detection_missed_and_remove_a_shipment(client, org, user):
    a, _, fr = make_shipment(org, two_products(), charges=[("Ocean Freight", "999.99")], boxes=1)
    b, _, _ = make_shipment(org, two_products(), charges=(), boxes=2)
    c, _, _ = make_shipment(org, two_products(), charges=(), boxes=1)
    client.force_login(user)
    client.post(reverse("landed:split_start", args=[a.pk]), {"invoice": fr.pk, "add": [b.pk, c.pk],
                                                              "basis": "containers"})
    assert [r.amount for r in allocation.allocations(fr)] == [D("250.00"), D("499.99"), D("250.00")]   # .75 .5 .75
    client.post(reverse("landed:split_save", args=[fr.pk]), {"basis": "equal", "remove": [c.pk]})
    assert [r.amount for r in allocation.allocations(fr)] == [D("500.00"), D("499.99")]
    r = client.post(reverse("landed:split_save", args=[fr.pk]), {"basis": "equal", "remove": [a.pk]}, follow=True)
    assert "be removed: the invoice is in that shipment" in r.content.decode()


def test_split_that_no_longer_adds_up_is_an_error_on_every_shipment(org, user):
    a, _, fr = make_shipment(org, two_products(), charges=[("Ocean Freight", "1000.00")])
    b, _, _ = make_shipment(org, two_products(), charges=())
    allocation.save_split(fr, user, "manual", amounts={a.pk: "600.00", b.pk: "400.00"}, add=[b.pk])
    correct_field(fr, "total_amount", "1100.00", user)
    after_correction(fr, "total_amount")
    for s in (a, b):   # the invoice's own shipment is checked by the correction, the other one by the split
        issue = ValidationIssue.objects.get(shipment=s, code="shared_split_mismatch", resolved=False)
        assert "add up to USD 1,000.00" in issue.message and issue.severity == "error"
    assert allocation.bill_lines(fr) is None
    assert "don't add up" in posting_blockers(a)[0]


def test_secondary_shipment_sees_errors_still_open_on_the_invoice(client, org, user):
    a, _, fr = make_shipment(org, two_products(), charges=[("Ocean Freight", "1000.00")])
    b, _, _ = make_shipment(org, two_products(), charges=())
    allocation.save_split(fr, user, "equal", add=[b.pk])
    ValidationIssue.objects.create(organization=org, shipment=a, document=fr, code="total_mismatch", severity="error",
                                   message="x", fingerprint="tm")
    client.force_login(user)
    page = client.get(reverse("review:shipment", args=[b.pk])).content.decode()
    assert "The invoice has 1 open error on" in page


def test_approval_limit_counts_the_shares_a_shipment_carries(org, user):
    a, _, fr = make_shipment(org, two_products(), charges=[("Ocean Freight", "1000.00")])
    b, _, _ = make_shipment(org, two_products(), charges=[("Trucking", "100.00")])
    allocation.save_split(fr, user, "manual", amounts={a.pk: "200.00", b.pk: "800.00"}, add=[b.pk])
    limited = get_user_model().objects.create_user("limited", password=PASSWORD)
    Membership.objects.create(user=limited, organization=org, role="approver", approval_limit=D("4500.00"))
    reasons = approval_blockers(b, limited)
    assert any("With its share of shared invoices, this shipment's total is USD 4,900.00" in x for x in reasons)


def test_landed_cost_counts_each_shipments_share_once(shared):
    doc, ships = shared["doc"], shared["ships"]
    total = D(shared["sc"]["invoice"]["total_amount"])
    charged = [landed.compute(Shipment.objects.get(pk=s.pk)).charges_total for s in ships]
    assert charged == [r.amount for r in allocation.allocations(doc)] and sum(charged) == total
    lc = landed.compute(ships[1])
    assert all(r.charge.shared for r in lc.charges)
    assert any(r.charge.category == "handling" and r.charge.description.startswith("Demurrage") for r in lc.charges)


def test_shared_invoice_permissions(client, shared, viewer, user):
    doc = shared["doc"]
    client.force_login(viewer)
    for name in ("split_save", "split_confirm", "split_dismiss", "split_reset"):
        assert client.post(reverse(f"landed:{name}", args=[doc.pk]), {"basis": "equal", "note": "abcdef"}
                           ).status_code == 403
    assert client.post(reverse("landed:split_start", args=[shared["ships"][0].pk])).status_code == 403
    page = client.get(reverse("review:shipment", args=[shared["ships"][0].pk])).content.decode()
    assert "Shared invoices" in page and "Confirm split" not in page
    client.force_login(user)
    page = client.get(reverse("review:shipment", args=[shared["ships"][0].pk])).content.decode()
    assert "Confirm split" in page and "Landed cost" in page and "Stainless cookware set 10pc" in page


# --------------------------------------------------------------------------- extraction and synthetic data


def test_rules_read_sku_hs_weight_and_volume_from_commercial_invoices():
    from apps.documents.services.extract_rules import extract_rules
    from apps.documents.services.ocr import read_text

    lines = synth.goods({"APX-CW10": 400, "APX-SU05": 800})
    pdf, truth = synth.commercial_invoice("CI-1", "PO-1", [make_container("OSL", 123456)], lines)
    got = extract_rules("commercial_invoice", read_text(pdf).text).values["line_items"]
    assert got == truth["line_items"]
    # Weights in a product name are per unit, so they are not read as the line's weight.
    from apps.documents.services.extract_rules import _line_items

    items = _line_items(["Description Qty Unit Price Amount", "Cement 50 kg bag 200 4.50 900.00"], goods_details=True)
    assert "weight_kg" not in items[0]
    plain = _line_items(["Description Qty Unit Price Amount", "Wall art canvas 10 4.00 40.00"], goods_details=True)
    assert plain == [{"description": "Wall art canvas", "quantity": "10", "unit_price": "4.00", "amount": "40.00"}]


def test_default_dataset_commercial_invoices_are_read_as_before(dataset):
    import json

    from apps.documents.services.extract_rules import extract_rules
    from apps.documents.services.ocr import read_text

    truth = json.loads((dataset / "ground_truth.json").read_text())
    ci = next(d for d in truth["documents"] if d["doc_type"] == "commercial_invoice" and not d["scanned"])
    got = extract_rules("commercial_invoice", read_text((dataset / "pdf" / ci["file"]).read_bytes()).text)
    assert all(set(li) <= {"description", "quantity", "unit_price", "amount"} for li in got.values["line_items"])


def test_schema_keeps_the_new_goods_fields():
    from apps.documents.schemas import lenient_model, wire_schema

    parsed = lenient_model("commercial_invoice").model_validate(
        {"line_items": [{"description": "Mug", "amount": "2", "sku": "M1", "hs_code": "6912.00", "weight_kg": 0.38}]})
    assert parsed.model_dump(mode="json", exclude_none=True)["line_items"][0] == {
        "description": "Mug", "amount": "2", "sku": "M1", "hs_code": "6912.00", "weight_kg": "0.38"}
    items = wire_schema("commercial_invoice")["properties"]["line_items"]["items"]
    assert items["additionalProperties"] is False and {"sku", "hs_code", "weight_kg", "volume_cbm"} <= set(
        items["required"])
    assert "sku" not in wire_schema("freight_invoice")["properties"]["line_items"]["items"]["properties"]


def test_synthetic_helpers_are_opt_in():
    import inspect

    import synthetic.generator as generator

    assert "landed" not in inspect.getsource(generator)
    sc = synth.scenario(seed=5, legs=2)
    assert len(sc["files"]) == 5 and sc["files"][-1][0].endswith("shared_freight_invoice.pdf")
    assert len(sc["invoice"]["legs"]) == 2


def test_seed_landed_command(org):
    from django.core.management import call_command

    call_command("seed_landed", "--org", org.slug, "--rounds", "1")
    assert SharedInvoice.objects.filter(organization=org).count() == 1
    call_command("seed_landed", "--org", org.slug, "--rounds", "1")   # idempotent
    assert Document.objects.filter(organization=org).count() == 7


def test_moving_documents_that_empties_a_shipment_repairs_the_split(client, shared, user):
    doc, (primary, second, third) = shared["doc"], shared["ships"]
    client.force_login(user)
    for d in list(third.documents):
        client.post(reverse("review:move_document", args=[d.pk]), {"target": second.pk})
    assert not Shipment.objects.filter(pk=third.pk).exists()
    rows = allocation.allocations(doc)
    assert {r.shipment_id for r in rows} == {primary.pk, second.pk}
    assert sum(r.amount for r in rows) == D(shared["sc"]["invoice"]["total_amount"])
    assert allocation.adds_up(doc, rows)


def test_a_failure_in_the_split_never_stops_matching(org, monkeypatch):
    def boom(*args, **kwargs):
        raise RuntimeError("bug")

    monkeypatch.setattr(allocation, "sync_related", boom)
    sc = synth.scenario(seed=55, legs=2)
    ingest_all(org, sc["files"])
    assert Document.objects.filter(organization=org, status="matched").count() == 5
    assert Shipment.objects.filter(organization=org).count() == 2


def test_detection_can_be_switched_off(org, settings):
    settings.SHARED_INVOICE_DETECT = False
    ingest_all(org, synth.scenario(seed=66, legs=2)["files"])
    assert not SharedInvoice.objects.filter(organization=org).exists()


# --------------------------------------------------------------------------- QA-003 / QA-024 / QA-025: what the page says


def _open_container_error(shipment):
    return ValidationIssue.objects.create(
        organization=shipment.organization, shipment=shipment, code="container_not_on_bl", severity="error",
        message="invoice.pdf: container ABCD1234567 is not on the bill of lading", fingerprint=f"c-{shipment.pk}")


def test_container_errors_on_a_shared_invoice_say_the_split_clears_them(client, approver, shared):
    doc, ships = shared["doc"], shared["ships"]
    _open_container_error(ships[0])
    number = doc.field("invoice_number")
    client.force_login(approver)

    page = client.get(reverse("review:shipment", args=[ships[0].pk])).content.decode()
    assert "container errors come from shared invoice" in page and number in page
    assert "they clear by themselves" in page

    client.post(reverse("landed:split_confirm", args=[doc.pk]))
    page = client.get(reverse("review:shipment", args=[ships[0].pk])).content.decode()
    assert "container errors come from shared invoice" not in page


def test_no_hint_when_the_container_error_has_nothing_to_do_with_a_shared_invoice(client, approver, org):
    s, _, _ = make_shipment(org, two_products())
    _open_container_error(s)
    client.force_login(approver)
    page = client.get(reverse("review:shipment", args=[s.pk])).content.decode()
    assert "container errors come from shared invoice" not in page


def test_a_ready_shipment_that_cannot_be_approved_is_not_labelled_ready(client, approver, shared):
    ships = shared["ships"]
    for s in ships:
        s.issues.all().update(resolved=True)
        Shipment.objects.filter(pk=s.pk).update(status="ready")
    client.force_login(approver)

    page = client.get(reverse("review:shipment", args=[ships[0].pk])).content.decode()

    assert "You can&#x27;t approve this yet" in page or "You can't approve this yet" in page
    assert "Waiting on a step" in page


def test_the_approve_button_asks_for_confirmation_like_the_shortcut(client, approver, approver2, shared):
    doc, ships = shared["doc"], shared["ships"]
    for s in ships:
        s.issues.all().update(resolved=True)
        Shipment.objects.filter(pk=s.pk).update(status="ready")
    client.force_login(approver2)
    client.post(reverse("landed:split_confirm", args=[doc.pk]))
    client.force_login(approver)
    page = client.get(reverse("review:shipment", args=[ships[1].pk])).content.decode()

    assert 'data-wf-confirm="wf-approve-dialog"' in page   # the button opens the same dialog as the "a" key
    assert 'id="wf-approve-dialog"' in page
    assert "Waiting on a step" not in page
