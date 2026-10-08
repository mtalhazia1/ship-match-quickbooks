"""Messy inputs for the intake tests: credit notes, multi-invoice batches, multi-page invoices,
spreadsheets, phone photos, multi-page TIFF scans and ZIP archives.

Only the tests use these helpers; the default dataset written by `generator.generate` does not change.
Every helper returns the file's bytes and, where useful, the values printed on it (the ground truth).
"""
from __future__ import annotations

import csv
import io
import zipfile
from datetime import date, timedelta
from decimal import Decimal

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from apps.shipments.services.containers import make_container
from synthetic.generator import CONSIGNEE, FORWARDERS, Party, money, render_freight_invoice

HARBORLINK = FORWARDERS[0]


# --------------------------------------------------------------------------- PDFs


def freight_invoice(number: str, bl: str, containers: list[str], *, vendor: Party = HARBORLINK, po: str = "PO-2026-7001",
                    day: date = date(2026, 4, 1), charges: list[tuple[str, str]] | None = None, variant: int = 0,
                    currency: str = "USD") -> tuple[bytes, dict]:
    """A one-page freight invoice in one of the generator's three layouts."""
    charges = charges or [("Ocean Freight", "2150.00"), ("Terminal Handling Charge (THC)", "310.00"),
                          ("Documentation Fee", "75.00")]
    items = [{"description": d, "amount": money(a)} for d, a in charges]
    total = sum((i["amount"] for i in items), Decimal("0.00"))
    d = {"vendor": vendor, "invoice_number": number, "invoice_date": day, "due_date": day + timedelta(days=30),
         "bl_number": bl, "container_numbers": containers, "po_numbers": [po], "currency": currency,
         "line_items": items, "printed_total": total}
    truth = {"vendor_name": vendor.name, "invoice_number": number, "invoice_date": day.isoformat(),
             "bl_number": bl, "container_numbers": containers, "po_numbers": [po], "currency": currency,
             "total_amount": f"{total:.2f}"}
    return render_freight_invoice(d, variant), truth


def credit_note(number: str, original_invoice: str | None, *, vendor: Party = HARBORLINK, bl: str | None = None,
                containers: list[str] | None = None, charges: list[tuple[str, str]] | None = None,
                style: str = "minus", day: date = date(2026, 4, 20), currency: str = "USD") -> tuple[bytes, dict]:
    """A vendor credit note. style: minus (-185.00), brackets ((185.00)) or plain (185.00)."""
    charges = charges or [("Terminal Handling Charge (THC) overcharge", "120.00"), ("Documentation Fee refund", "65.00")]
    total = sum((Decimal(a) for _, a in charges), Decimal("0.00"))

    def shown(amount) -> str:
        value = f"{Decimal(amount):,.2f}"
        return {"minus": f"-{value}", "brackets": f"({value})"}.get(style, value)

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    y = 800

    def line(*parts, size=10, bold=False, step=15):
        nonlocal y
        for p in parts:
            c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
            (c.drawRightString if len(p) > 2 else c.drawString)(p[0], y, p[1])
        y -= step

    line((50, vendor.name), (545, "CREDIT NOTE", "r"), size=13, bold=True, step=16)
    line((50, vendor.address), size=9)
    y -= 10
    line((50, "Credit Note No.:"), (190, number))
    line((50, "Date:"), (190, day.isoformat()))
    if original_invoice:
        line((50, "Original Invoice No.:"), (190, original_invoice))
    if bl:
        line((50, "B/L No.:"), (190, bl))
    if containers:
        line((50, "Container(s):"), (190, ", ".join(containers)))
    line((50, "Currency:"), (190, currency))
    line((50, "Bill To:"), (190, CONSIGNEE.name))
    y -= 14
    line((50, "Description"), (545, f"Amount ({currency})", "r"), bold=True)
    for desc, amount in charges:
        line((50, desc), (545, shown(amount), "r"))
    y -= 6
    line((330, f"Total Credit ({currency}):"), (545, shown(total), "r"), bold=True)
    y -= 20
    line((50, "This credit will be applied to your account. No payment is due."), size=8)
    c.showPage()
    c.save()
    truth = {"vendor_name": vendor.name, "credit_note_number": number, "original_invoice_number": original_invoice,
             "invoice_date": day.isoformat(), "currency": currency, "total_amount": f"{total:.2f}",
             "line_items": [{"description": d, "amount": f"{Decimal(a):.2f}"} for d, a in charges]}
    return buf.getvalue(), truth


def concatenate(pdfs: list[bytes]) -> bytes:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for pdf in pdfs:
        for page in PdfReader(io.BytesIO(pdf)).pages:
            writer.add_page(page)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def batch(count: int = 3, owner: str = "OSL", bl_prefix: str = "OSLN") -> tuple[bytes, list[dict]]:
    """A carrier batch: `count` one-page invoices, each for its own shipment, in one PDF."""
    pdfs, truths = [], []
    for n in range(count):
        bl = f"{bl_prefix}{7700000100 + n * 11}"
        containers = [make_container(owner, 640100 + n * 7)]
        pdf, truth = freight_invoice(f"HA-90{n:04d}", bl, containers, po=f"PO-2026-8{n:03d}", variant=n % 3,
                                     day=date(2026, 5, 1) + timedelta(days=n))
        pdfs.append(pdf)
        truths.append(truth)
    return concatenate(pdfs), truths


def multipage_invoice(pages: int = 3, page_numbers: bool = True, repeat_header: bool = False,
                      lines_per_page: int = 30) -> tuple[bytes, dict]:
    """One freight invoice whose charges run over several pages; only the last page has the total."""
    number, bl = "HA-550077", "OSLN7712340001"
    container = make_container("OSL", 712340)
    charges = [(f"Detention day {i + 1} {container}", Decimal("45.00") + i) for i in range(pages * lines_per_page - 5)]
    total = sum((a for _, a in charges), Decimal("0.00"))
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    per_page = [charges[i:i + lines_per_page] for i in range(0, len(charges), lines_per_page)]
    for p, chunk in enumerate(per_page):
        y = 800
        c.setFont("Helvetica-Bold", 13)
        if p == 0 or repeat_header:
            c.drawString(50, y, HARBORLINK.name)
            c.drawRightString(545, y, "FREIGHT INVOICE")
            y -= 18
            c.setFont("Helvetica", 10)
            c.drawString(50, y, f"Invoice No.: {number}")
            y -= 15
        if p == 0:
            c.setFont("Helvetica", 10)
            for label, value in (("Invoice Date:", "2026-05-03"), ("B/L No.:", bl), ("Container(s):", container),
                                 ("Customer Ref / PO:", "PO-2026-9001"), ("Bill To:", CONSIGNEE.name)):
                c.drawString(50, y, label)
                c.drawString(170, y, value)
                y -= 15
            y -= 10
            c.setFont("Helvetica-Bold", 10)
            c.drawString(50, y, "Charge Description")
            c.drawRightString(545, y, "Amount (USD)")
            y -= 15
        c.setFont("Helvetica", 10)
        for desc, amount in chunk:
            c.drawString(50, y, desc)
            c.drawRightString(545, y, f"{amount:,.2f}")
            y -= 14
        if p == len(per_page) - 1:
            y -= 6
            c.setFont("Helvetica-Bold", 10)
            c.drawString(330, y, "TOTAL DUE:")
            c.drawRightString(545, y, f"{total:,.2f}")
        if page_numbers:
            c.setFont("Helvetica", 8)
            c.drawRightString(545, 30, f"Page {p + 1} of {len(per_page)}")
        c.showPage()
    c.save()
    return buf.getvalue(), {"invoice_number": number, "bl_number": bl, "total_amount": f"{total:.2f}",
                            "pages": len(per_page)}


# --------------------------------------------------------------------------- spreadsheets


def invoice_xlsx(bl: str, containers: list[str], *, number: str = "HL-552310", po: str = "PO-2026-1001",
                 vendor: Party = HARBORLINK, currency: str = "USD") -> tuple[bytes, dict]:
    """A carrier-style Excel invoice: a header block of label/value cells, a charge table and a total row."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    rate = [("Ocean Freight", 2, Decimal("2150.00")), ("Terminal Handling Charge (THC)", 2, Decimal("310.00")),
            ("Documentation Fee", 1, Decimal("75.00"))]
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["A1"] = vendor.name
    ws["A1"].font = Font(bold=True, size=14)
    ws["E1"] = "FREIGHT INVOICE"
    ws["A2"] = vendor.address
    ws["A4"], ws["B4"] = "Invoice No.", number
    ws["D4"], ws["E4"] = "Invoice Date", date(2026, 3, 28)
    ws["A5"], ws["B5"] = "B/L No.", bl
    ws["D5"], ws["E5"] = "Due Date", date(2026, 4, 27)
    ws["A6"], ws["B6"] = "Currency", currency
    ws["D6"], ws["E6"] = "Customer Ref / PO", po
    for col, head in zip("ABCDE", ["Container", "Description", "Qty", "Rate", "Amount"]):
        ws[f"{col}8"] = head
        ws[f"{col}8"].font = Font(bold=True)
    row, total, items = 9, Decimal("0.00"), []
    for desc, qty, price in rate:
        amount = price * qty
        ws[f"A{row}"] = ", ".join(containers) if qty > 1 else containers[0]
        ws[f"B{row}"], ws[f"C{row}"], ws[f"D{row}"], ws[f"E{row}"] = desc, qty, float(price), float(amount)
        ws[f"D{row}"].number_format = ws[f"E{row}"].number_format = "#,##0.00"
        items.append({"description": desc, "quantity": str(qty), "unit_price": f"{price:.2f}", "amount": f"{amount:.2f}"})
        total += amount
        row += 1
    ws[f"D{row + 1}"], ws[f"E{row + 1}"] = "Total due", float(total)
    ws[f"E{row + 1}"].number_format = "#,##0.00"
    ws[f"A{row + 3}"] = "Please remit within 30 days. Bank details on request."
    out = io.BytesIO()
    wb.save(out)
    truth = {"vendor_name": vendor.name, "invoice_number": number, "invoice_date": "2026-03-28",
             "due_date": "2026-04-27", "bl_number": bl, "container_numbers": containers, "po_numbers": [po],
             "currency": currency, "total_amount": f"{total:.2f}", "line_items": items}
    return out.getvalue(), truth


def invoice_csv(bl: str, container: str, *, number: str = "AT-946721",
                vendor: Party = FORWARDERS[2]) -> tuple[bytes, dict]:
    """A billing-system export: one row per charge with the invoice's references on every row."""
    charges = [("Ocean Freight", "4668.02"), ("Terminal Handling Charge (THC)", "701.18")]
    total = sum((Decimal(a) for _, a in charges), Decimal("0.00"))
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["Invoice Number", "Invoice Date", "Vendor", "B/L Number", "Container", "Charge Description",
                "Amount (USD)"])
    for desc, amount in charges:
        w.writerow([number, "05/10/2026", vendor.name, bl, container, desc, amount])
    w.writerow(["", "", "", "", "", "Total", f"{total:.2f}"])
    truth = {"vendor_name": vendor.name, "invoice_number": number, "invoice_date": "2026-05-10", "bl_number": bl,
             "container_numbers": [container], "currency": "USD", "total_amount": f"{total:.2f}",
             "line_items": [{"description": d, "amount": a} for d, a in charges]}
    return out.getvalue().encode("utf-8"), truth


# --------------------------------------------------------------------------- images


def page_image(pdf: bytes, page: int = 0, resolution: int = 110):
    """A PDF page as a Pillow image, like a phone photo of a printed invoice (grayish, slightly tilted)."""
    import pdfplumber
    from PIL import Image

    with pdfplumber.open(io.BytesIO(pdf)) as doc:
        img = doc.pages[page].to_image(resolution=resolution).original.convert("RGB")
    img = img.rotate(0.8, expand=True, fillcolor=(236, 236, 230), resample=Image.Resampling.BICUBIC)
    return img


def photo(pdf: bytes, fmt: str = "PNG", sideways: bool = False) -> bytes:
    """A photo of the first page. sideways=True stores the pixels turned 90 degrees with an EXIF
    orientation tag saying how to turn them back, as phone cameras do."""
    img = page_image(pdf)
    out = io.BytesIO()
    if sideways:
        img = img.rotate(90, expand=True)  # pixels lie on their side...
        exif = img.getexif()
        exif[0x0112] = 6                   # ...and the camera says "turn 90 degrees clockwise to view"
        img.save(out, format=fmt, exif=exif.tobytes(), quality=85) if fmt == "JPEG" else img.save(out, format=fmt,
                                                                                                    exif=exif.tobytes())
    else:
        img.save(out, format=fmt, **({"quality": 85} if fmt == "JPEG" else {}))
    return out.getvalue()


def tiff_scan(pdfs: list[bytes]) -> bytes:
    """A multi-page TIFF (a scanner's output): one frame per PDF's first page."""
    frames = [page_image(pdf).convert("L") for pdf in pdfs]
    out = io.BytesIO()
    frames[0].save(out, format="TIFF", save_all=True, append_images=frames[1:], compression="tiff_deflate")
    return out.getvalue()


# --------------------------------------------------------------------------- archives


def zip_of(entries: dict[str, bytes], compression: int = zipfile.ZIP_DEFLATED) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=compression) as zf:
        for name, data in entries.items():
            # Fixed timestamp: the same entries always give the same bytes (and the same SHA-256)
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = compression
            zf.writestr(info, data)
    return out.getvalue()
